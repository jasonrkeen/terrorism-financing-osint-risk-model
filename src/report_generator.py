"""
Report generator module for the terrorism financing OSINT risk model.

This script creates a Word report and, when supported by the local system,
exports a PDF version of the report.

The report is designed for defensive screening, academic analysis, and
financial-crime risk research. It does not provide operational terrorism-
financing guidance and does not identify individuals, entities, transactions,
wallets, accounts, or networks.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from pydoc import doc
from typing import Iterable

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from src.config import (
    PROJECT_TITLE,
    STUDY_YEAR,
    RISK_RANKINGS_FILE,
    CORRELATION_MATRIX_FILE,
    RISK_TIER_SUMMARY_FILE,
    MODEL_PERFORMANCE_FILE,
    FALSE_POSITIVES_FILE,
    FALSE_NEGATIVES_FILE,
    SENSITIVITY_OVERLAP_FILE,
    CURRENT_DATA_SUPPLEMENT_FILE,
    CURRENT_FATF_STATUS_COMPARISON_FILE,
    CONFLICT_DATA_SUPPLEMENT_FILE,
    TOP_RISK_DRIVERS_FILE,
    REGIONAL_RISK_SUMMARY_FILE,
    ADDITIONAL_RISK_FACTOR_SUMMARY_FILE,
    SANCTIONS_EXPOSURE_SUPPLEMENT_FILE,
    FATF_MUTUAL_EVALUATION_SUPPLEMENT_FILE,
    LOGIT_SUMMARY_FILE,
    WORD_REPORT_FILE,
    PDF_REPORT_FILE,
    TOP_15_RISK_CHART_FILE,
    RISK_SCORE_PROBABILITY_CHART_FILE,
    GREY_LIST_PROBABILITY_TOP_20_CHART_FILE,
    RISK_TIER_DISTRIBUTION_CHART_FILE,
    SENSITIVITY_OVERLAP_CHART_FILE,
    TOP_15_CONFLICT_FATALITIES_CHART_FILE,
    REGIONAL_AVERAGE_RISK_CHART_FILE,
    REGIONAL_FATF_COUNT_CHART_FILE,
    CURRENT_FATF_STATUS_COMPARISON_CHART_FILE,
    TOP_15_COLLECTION_PRIORITY_CHART_FILE,
    INTELLIGENCE_GAP_DISTRIBUTION_CHART_FILE,
    ACLED_SIGNAL_TOP_15_CHART_FILE,
    ACLED_SIGNAL_ENRICHMENT_FILE,
    GDELT_NEWS_SIGNAL_ENRICHMENT_FILE,
    TOP_15_GDELT_NEWS_SIGNAL_CHART_FILE,
)


# ---------------------------------------------------
# Display settings
# ---------------------------------------------------

MAX_TOP_RISK_ROWS = 20
MAX_TOP_PROBABILITY_ROWS = 10
MAX_FALSE_POSITIVE_ROWS = 8
MAX_FALSE_NEGATIVE_ROWS = 8
MAX_TOP_DRIVER_ROWS = 10
MAX_CONFLICT_ROWS = 12
MAX_CURRENT_FATF_ROWS = 12
MAX_CELL_TEXT_LENGTH = 120
MAX_COLLECTION_PRIORITY_ROWS = 15
MAX_SOURCE_REGISTER_ROWS = 12
MAX_ACLED_SIGNAL_ROWS = 12
MAX_ORGANIZED_CRIME_ROWS = 12
MAX_GDELT_SIGNAL_ROWS = 12

THREAT_FINANCE_COLLECTION_PRIORITIES_FILE = Path(
    "outputs/tables/threat_finance_collection_priorities.csv"
)

THREAT_FINANCE_SOURCE_REGISTER_FILE = Path(
    "outputs/tables/source_register.csv"
)

ORGANIZED_CRIME_ENRICHMENT_FILE = Path(
    "outputs/tables/organized_crime_enrichment.csv"
)

TOP_15_ORGANIZED_CRIME_CHART_FILE = Path(
    "outputs/charts/top_15_organized_crime_vulnerability_scores.png"
)

REPORT_AUTHOR_NAME = "Jason Keen"
REPORT_AUTHOR_CONTACT = "keenjasonr@gmail.com"


INDICATOR_YEAR_COLUMNS = [
    "control_of_corruption_year",
    "rule_of_law_year",
    "political_stability_year",
    "gdp_per_capita_year",
    "inflation_year",
    "unemployment_year",
]


RISK_FEATURE_LABELS = {
    "corruption_risk_scaled": "Corruption Risk",
    "rule_of_law_risk_scaled": "Rule of Law Risk",
    "political_instability_risk_scaled": "Instability Risk",
    "low_income_risk_scaled": "Low Income Risk",
    "inflation_risk_scaled": "Inflation Risk",
    "unemployment_risk_scaled": "Unemployment Risk",
    "tf_vulnerability_score": "Risk Score",
    "fatf_grey_list": "Grey List",
    "predicted_grey_list_probability": "Pred. Prob.",
}


# ---------------------------------------------------
# General helpers
# ---------------------------------------------------

def read_csv_if_exists(path: Path) -> pd.DataFrame | None:
    """
    Read a CSV file if it exists. Return None when unavailable.
    """

    if not Path(path).exists():
        print(f"Optional file not found, skipping: {path}")
        return None

    try:
        return pd.read_csv(path)
    except Exception as error:
        print(f"Could not read {path}: {error}")
        return None


def read_text_if_exists(path: Path) -> str:
    """
    Read a text file if it exists. Return an empty string when unavailable.
    """

    if not Path(path).exists():
        print(f"Optional text file not found, skipping: {path}")
        return ""

    try:
        return Path(path).read_text(encoding="utf-8")
    except Exception as error:
        print(f"Could not read {path}: {error}")
        return ""


def safe_percent(value: float | int | str | None) -> str:
    """
    Format a numeric value as a percentage.
    """

    if pd.isna(value):
        return "N/A"

    try:
        return f"{float(value) * 100:.2f}%"
    except (TypeError, ValueError):
        return "N/A"


def safe_number(value: float | int | str | None, decimals: int = 2) -> str:
    """
    Format a numeric value safely.
    """

    if pd.isna(value):
        return "N/A"

    try:
        return f"{float(value):,.{decimals}f}"
    except (TypeError, ValueError):
        return str(value)


def safe_int(value: float | int | str | None) -> str:
    """
    Format a numeric value as an integer.
    """

    if pd.isna(value):
        return "N/A"

    try:
        return f"{int(round(float(value))):,}"
    except (TypeError, ValueError):
        return str(value)




def safe_year(value: float | int | str | None) -> str:
    """
    Format year values without thousands separators.

    This prevents years such as 2024 from displaying as 2,024 in report tables.
    """

    if pd.isna(value):
        return "N/A"

    try:
        return str(int(round(float(value))))
    except (TypeError, ValueError):
        return str(value)

def yes_no(value: float | int | str | bool | None) -> str:
    """
    Convert a binary flag to Yes/No.
    """

    if pd.isna(value):
        return "N/A"

    try:
        return "Yes" if int(value) == 1 else "No"
    except (TypeError, ValueError):
        return str(value)




def pluralize(count: int, singular: str, plural: str | None = None) -> str:
    """
    Return a correctly pluralized noun phrase.
    """

    if plural is None:
        plural = f"{singular}s"

    return singular if count == 1 else plural

def add_paragraphs(doc: Document, paragraphs: Iterable[str]) -> None:
    """
    Add multiple paragraphs to the document.
    """

    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)


def set_document_style(doc: Document) -> None:
    """
    Set basic report styling.
    """

    styles = doc.styles

    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)

    for style_name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        if style_name in styles:
            styles[style_name].font.name = "Aptos"

    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)


def add_title_page(doc: Document) -> None:
    """
    Add the report title page.
    """

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(PROJECT_TITLE)
    title_run.bold = True
    title_run.font.size = Pt(20)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(
        "Defensive Open-Source Intelligence and Financial Risk Analysis"
    )
    subtitle_run.italic = True
    subtitle_run.font.size = Pt(13)

    doc.add_paragraph()

    study_year = doc.add_paragraph()
    study_year.alignment = WD_ALIGN_PARAGRAPH.CENTER
    study_year.add_run(f"Target Study Year: {STUDY_YEAR}").bold = True

    generated = doc.add_paragraph()
    generated.alignment = WD_ALIGN_PARAGRAPH.CENTER
    generated.add_run(f"Generated: {datetime.now():%B %d, %Y}")

    author = doc.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author_run = author.add_run(
        f"Prepared by: {REPORT_AUTHOR_NAME} | {REPORT_AUTHOR_CONTACT}"
    )
    author_run.font.size = Pt(10)

    doc.add_paragraph()

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.add_run(
        "This report is intended for defensive screening, academic analysis, "
        "and financial-crime risk research."
    ).italic = True

    doc.add_page_break()


def shorten_output_path(value: object) -> str:
    """
    Convert long local Windows paths into short project-relative paths.

    This keeps generated report tables readable while still showing where
    outputs are saved inside the project.
    """

    if pd.isna(value):
        return "N/A"

    text = str(value)
    normalized = text.replace("\\", "/")

    markers = [
        "outputs/tables/",
        "outputs/charts/",
        "outputs/reports/",
        "data/raw/",
        "data/processed/",
    ]

    for marker in markers:
        marker_index = normalized.find(marker)
        if marker_index >= 0:
            return normalized[marker_index:]

    return text


def clean_table_value(value: object, max_length: int = MAX_CELL_TEXT_LENGTH) -> str:
    """
    Clean and shorten values before writing them into Word tables.
    """

    text = shorten_output_path(value)
    text = text.replace("\n", " ").replace("\r", " ").strip()

    while "  " in text:
        text = text.replace("  ", " ")

    if len(text) > max_length:
        return text[: max_length - 3].rstrip() + "..."

    return text


def format_word_table(table, font_size: int = 8) -> None:
    """
    Apply compact formatting to Word tables.
    """

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)

                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(font_size)
                    if row_index == 0:
                        run.bold = True


def add_dataframe_table(
    doc: Document,
    df: pd.DataFrame | None,
    columns: list[str],
    headers: list[str],
    max_rows: int | None = None,
) -> None:
    """
    Add a pandas DataFrame as a Word table.
    """

    if df is None or df.empty:
        doc.add_paragraph("No data available.")
        return

    display_df = df.copy()

    if max_rows is not None:
        display_df = display_df.head(max_rows)

    available_columns = [column for column in columns if column in display_df.columns]

    if not available_columns:
        doc.add_paragraph("No display columns available.")
        return

    headers = headers[: len(available_columns)]

    table = doc.add_table(rows=1, cols=len(available_columns))
    table.style = "Table Grid"

    header_cells = table.rows[0].cells

    for idx, header in enumerate(headers):
        header_cells[idx].text = clean_table_value(header, max_length=60)

    for _, row in display_df.iterrows():
        cells = table.add_row().cells

        for idx, column in enumerate(available_columns):
            value = row[column]
            cells[idx].text = clean_table_value(value)

    format_word_table(table)
    doc.add_paragraph()


def add_picture_if_exists(
    doc: Document,
    image_path: Path,
    caption: str,
    width: float = 6.2,
) -> None:
    """
    Add an image and caption if the image exists.
    """

    if not Path(image_path).exists():
        print(f"Chart not found, skipping: {image_path}")
        return

    doc.add_paragraph(caption).bold = True

    try:
        doc.add_picture(str(image_path), width=Inches(width))
    except Exception as error:
        doc.add_paragraph(f"Could not insert chart: {image_path}")
        print(f"Could not insert chart {image_path}: {error}")

    doc.add_paragraph()


# ---------------------------------------------------
# Data preparation helpers
# ---------------------------------------------------

def prepare_top_rankings(df: pd.DataFrame) -> pd.DataFrame:
    """
    Prepare the top risk ranking table.
    """

    output = df.sort_values("tf_vulnerability_score", ascending=False).copy()
    output["Score"] = output["tf_vulnerability_score"].apply(
        lambda value: safe_number(value, 2)
    )
    output["Grey List"] = output["fatf_grey_list"].apply(yes_no)

    if "predicted_grey_list_probability" in output.columns:
        output["Pred. Prob."] = output["predicted_grey_list_probability"].apply(
            safe_percent
        )
    else:
        output["Pred. Prob."] = "N/A"

    return output.rename(
        columns={
            "risk_rank": "Rank",
            "country": "Country",
        }
    )


def prepare_top_probabilities(df: pd.DataFrame) -> pd.DataFrame:
    """
    Prepare top predicted grey-list probability table.
    """

    if "predicted_grey_list_probability" not in df.columns:
        return pd.DataFrame()

    output = df.sort_values(
        "predicted_grey_list_probability",
        ascending=False,
    ).copy()

    output["Risk Score"] = output["tf_vulnerability_score"].apply(
        lambda value: safe_number(value, 2)
    )
    output["Actual Grey List"] = output["fatf_grey_list"].apply(yes_no)
    output["Pred. Probability"] = output[
        "predicted_grey_list_probability"
    ].apply(safe_percent)

    return output.rename(columns={"country": "Country"})


def prepare_risk_tier_summary(df: pd.DataFrame) -> pd.DataFrame:
    """
    Prepare risk tier summary table.
    """

    output = df.copy()

    if "average_score" in output.columns:
        output["Avg. Score"] = output["average_score"].apply(
            lambda value: safe_number(value, 2)
        )

    if "average_predicted_probability" in output.columns:
        output["Avg. Probability"] = output[
            "average_predicted_probability"
        ].apply(safe_percent)

    if "fatf_grey_list_count" in output.columns:
        output["FATF Count"] = output["fatf_grey_list_count"].apply(safe_int)

    if "country_count" in output.columns:
        output["Countries"] = output["country_count"].apply(safe_int)

    return output.rename(columns={"risk_tier": "Risk Tier"})


def format_performance_value(value: float | int | str | None) -> str:
    """
    Safely format model-performance values.
    """

    if pd.isna(value):
        return "N/A"

    try:
        numeric_value = float(value)
    except (TypeError, ValueError):
        return str(value)

    if numeric_value < 10:
        return safe_number(numeric_value, 3)

    return safe_int(numeric_value)


def prepare_model_performance(df: pd.DataFrame) -> pd.DataFrame:
    """
    Prepare model performance summary table.
    """

    labels = {
        "accuracy": "Accuracy",
        "precision_class_1": "Precision, Grey-List Class",
        "recall_class_1": "Recall, Grey-List Class",
        "f1_class_1": "F1 Score, Grey-List Class",
        "roc_auc": "ROC AUC",
        "test_set_size": "Test Set Size",
        "positive_cases_in_test_set": "Grey-List Cases in Test Set",
        "negative_cases_in_test_set": "Non-Grey-List Cases in Test Set",
        "positive_cases_total": "Grey-List Cases, Total",
        "negative_cases_total": "Non-Grey-List Cases, Total",
    }

    output = df.copy()
    output["Metric"] = output["metric"].map(labels).fillna(output["metric"])
    output["Value"] = output.apply(
        lambda row: safe_int(row["value"])
        if str(row["metric"]).endswith("_size")
        or "cases" in str(row["metric"])
        else format_performance_value(row["value"]),
        axis=1,
    )

    return output


def prepare_false_output(df: pd.DataFrame) -> pd.DataFrame:
    """
    Prepare false positive or false negative table.
    """

    output = df.copy()

    output["Score"] = output["tf_vulnerability_score"].apply(
        lambda value: safe_number(value, 2)
    )
    output["Pred. Prob."] = output["predicted_grey_list_probability"].apply(
        safe_percent
    )

    return output.rename(
        columns={
            "country": "Country",
            "risk_rank": "Rank",
        }
    )


def prepare_sensitivity_overlap(df: pd.DataFrame) -> pd.DataFrame:
    """
    Prepare sensitivity overlap summary table.
    """

    output = df.copy()
    output["Scenario"] = (
        output["scenario"].astype(str).str.replace("_", " ", regex=False).str.title()
    )
    output["Top-20 Overlap"] = output["top_20_overlap_with_baseline"].apply(
        safe_int
    )
    output["Baseline Count"] = output["baseline_top_20_count"].apply(safe_int)
    output["Overlap %"] = output["overlap_percentage"].apply(safe_percent)

    return output


def prepare_risk_drivers(df: pd.DataFrame) -> pd.DataFrame:
    """
    Prepare top risk drivers table.
    """

    output = df.copy()

    rename_map = {
        "risk_rank": "Rank",
        "country": "Country",
        "tf_vulnerability_score": "Score",
        "top_driver_1": "Driver 1",
        "top_driver_2": "Driver 2",
        "top_driver_3": "Driver 3",
        "driver_1": "Driver 1",
        "driver_2": "Driver 2",
        "driver_3": "Driver 3",
    }

    output = output.rename(columns=rename_map)

    if "Score" in output.columns:
        output["Score"] = output["Score"].apply(lambda value: safe_number(value, 2))

    return output


def prepare_conflict_table(df: pd.DataFrame) -> pd.DataFrame:
    """
    Prepare conflict supplement table.
    """

    output = df.copy()

    if "conflict_deaths_total" in output.columns:
        output = output.sort_values("conflict_deaths_total", ascending=False)

    output["Risk Score"] = output["tf_vulnerability_score"].apply(
        lambda value: safe_number(value, 2)
    )

    if "conflict_deaths_total" in output.columns:
        output["Deaths"] = output["conflict_deaths_total"].apply(safe_int)
    else:
        output["Deaths"] = "N/A"

    if "conflict_dyad_count" in output.columns:
        output["Conflict Events"] = output["conflict_dyad_count"].apply(safe_int)
    elif "conflict_event_count" in output.columns:
        output["Conflict Events"] = output["conflict_event_count"].apply(safe_int)
    else:
        output["Conflict Events"] = "N/A"

    if "ucdp_year" in output.columns:
        output["UCDP Year"] = output["ucdp_year"].apply(safe_year)
    else:
        output["UCDP Year"] = "N/A"

    return output.rename(
        columns={
            "country": "Country",
            "risk_rank": "Rank",
            "conflict_intensity_category": "Category",
        }
    )


def prepare_current_fatf_table(df: pd.DataFrame) -> pd.DataFrame:
    """
    Prepare current FATF monitored jurisdictions table.
    """

    output = df.copy()

    if "current_fatf_increased_monitoring" in output.columns:
        output = output[output["current_fatf_increased_monitoring"] == 1]

    if "predicted_grey_list_probability" in output.columns:
        output = output.sort_values(
            "predicted_grey_list_probability",
            ascending=False,
        )

    output["Score"] = output["tf_vulnerability_score"].apply(
        lambda value: safe_number(value, 2)
    )
    output["Pred. Prob."] = output["predicted_grey_list_probability"].apply(
        safe_percent
    )

    if "model_fatf_grey_list" in output.columns:
        output["Model FATF"] = output["model_fatf_grey_list"].apply(yes_no)
    elif "fatf_grey_list" in output.columns:
        output["Model FATF"] = output["fatf_grey_list"].apply(yes_no)
    elif "baseline_fatf_grey_list" in output.columns:
        output["Model FATF"] = output["baseline_fatf_grey_list"].apply(yes_no)

    if "current_fatf_increased_monitoring" in output.columns:
        output["Current FATF"] = output["current_fatf_increased_monitoring"].apply(
            yes_no
        )

    return output.rename(
        columns={
            "country": "Country",
            "risk_rank": "Rank",
        }
    )


def prepare_collection_priority_table(df: pd.DataFrame) -> pd.DataFrame:
    """
    Prepare threat finance collection priority table for the report.
    """

    output = df.copy()

    if "collection_priority_score" in output.columns:
        output = output.sort_values(
            "collection_priority_score",
            ascending=False,
        )

    rename_map = {
        "country": "Country",
        "iso3": "ISO3",
        "baseline_tf_risk_score": "Baseline Risk",
        "baseline_tf_risk_level": "Risk Level",
        "live_signal_score": "Live Signal",
        "source_staleness_score": "Source Staleness",
        "intelligence_gap_score": "Intel Gap",
        "intelligence_gap_level": "Gap Level",
        "collection_priority_score": "Collection Priority",
        "collection_priority_level": "Priority Level",
        "recommended_collection_focus": "Recommended Focus",
        "analyst_note": "Analyst Note",
    }

    output = output.rename(columns=rename_map)

    numeric_columns = [
        "Baseline Risk",
        "Live Signal",
        "Source Staleness",
        "Intel Gap",
        "Collection Priority",
    ]

    for column in numeric_columns:
        if column in output.columns:
            output[column] = output[column].apply(lambda value: safe_number(value, 2))

    return output


def prepare_acled_signal_table(df: pd.DataFrame) -> pd.DataFrame:
    """
    Prepare ACLED recent conflict and unrest signal table for the report.
    """

    output = df.copy()

    if "acled_conflict_signal_score" in output.columns:
        output = output.sort_values(
            "acled_conflict_signal_score",
            ascending=False,
        )

    rename_map = {
        "country": "Country",
        "iso3": "ISO / Code",
        "acled_event_count": "Events",
        "acled_fatalities": "Fatalities",
        "acled_political_violence_events": "Political Violence",
        "acled_unrest_events": "Unrest",
        "acled_conflict_signal_score": "ACLED Signal",
        "acled_signal_level": "Signal Level",
        "acled_latest_event_date": "Latest Event",
    }

    output = output.rename(columns=rename_map)

    for column in ["Events", "Fatalities", "Political Violence", "Unrest"]:
        if column in output.columns:
            output[column] = output[column].apply(safe_int)

    if "ACLED Signal" in output.columns:
        output["ACLED Signal"] = output["ACLED Signal"].apply(
            lambda value: safe_number(value, 2)
        )

    return output


def prepare_organized_crime_table(df: pd.DataFrame) -> pd.DataFrame:
    """
    Prepare Global Organized Crime Index enrichment table for the report.

    The enrichment output may contain both country_code and iso3. This function
    keeps one ISO3 display field to avoid duplicate column names in the Word
    report table.
    """

    output = df.copy()

    if "organized_crime_score" in output.columns:
        output = output.sort_values(
            "organized_crime_score",
            ascending=False,
        )

    if "country_code" in output.columns:
        output["display_iso3"] = output["country_code"]
    elif "iso3" in output.columns:
        output["display_iso3"] = output["iso3"]
    else:
        output["display_iso3"] = "N/A"

    rename_map = {
        "country": "Country",
        "display_iso3": "ISO3",
        "oc_index_year": "Index Year",
        "criminality_score": "Criminality",
        "criminal_markets_score": "Markets",
        "criminal_actors_score": "Actors",
        "resilience_score": "Resilience",
        "organized_crime_exposure_score": "Exposure",
        "organized_crime_resilience_gap_score": "Resilience Gap",
        "organized_crime_score": "OC Vulnerability",
        "organized_crime_level": "OC Level",
    }

    output = output.rename(columns=rename_map)

    # Drop duplicate column names defensively before Word table rendering.
    output = output.loc[:, ~output.columns.duplicated()].copy()

    numeric_columns = [
        "Criminality",
        "Markets",
        "Actors",
        "Resilience",
        "Exposure",
        "Resilience Gap",
        "OC Vulnerability",
    ]

    for column in numeric_columns:
        if column in output.columns:
            output[column] = output[column].apply(lambda value: safe_number(value, 2))

    if "Index Year" in output.columns:
        output["Index Year"] = output["Index Year"].apply(safe_year)

    return output


def prepare_gdelt_signal_table(df: pd.DataFrame) -> pd.DataFrame:
    """
    Prepare GDELT 2.0 batch public-news signal enrichment table for the report.

    The current GDELT layer uses batch GKG files. Some backward-compatible
    columns from the earlier DOC API prototype may still appear in outputs, but
    the preferred display fields are batch-file count, matched GKG records,
    matched source count, average tone, and signal score.
    """

    output = df.copy()

    if "gdelt_news_signal_score" in output.columns:
        output = output.sort_values(
            "gdelt_news_signal_score",
            ascending=False,
        )

    rename_map = {
        "country": "Country",
        "country_code": "ISO3",
        "gdelt_batch_file_count": "Batch Files",
        "gdelt_matching_record_count": "Matched Records",
        "gdelt_matching_source_count": "Matched Sources",
        "gdelt_average_tone": "Average Tone",
        "gdelt_record_count_score": "Record Score",
        "gdelt_source_diversity_score": "Source Score",
        "gdelt_negative_tone_score": "Tone Score",
        "gdelt_news_signal_score": "GDELT Signal",
        "gdelt_news_signal_level": "Signal Level",
        "gdelt_query_status": "Batch Status",
        "gdelt_confidence": "Confidence",
        "gdelt_days_since_update": "Days Since Update",
        "gdelt_latest_file_datetime": "Latest GDELT File",
    }

    output = output.rename(columns=rename_map)

    integer_columns = [
        "Batch Files",
        "Matched Records",
        "Matched Sources",
        "Days Since Update",
    ]

    decimal_columns = [
        "Average Tone",
        "Record Score",
        "Source Score",
        "Tone Score",
        "GDELT Signal",
    ]

    for column in integer_columns:
        if column in output.columns:
            output[column] = output[column].apply(safe_int)

    for column in decimal_columns:
        if column in output.columns:
            output[column] = output[column].apply(lambda value: safe_number(value, 2))

    return output


def prepare_source_register_table(df: pd.DataFrame) -> pd.DataFrame:
    """
    Prepare source register table for the report.
    """

    output = df.copy()

    rename_map = {
        "source_name": "Source",
        "source_type": "Type",
        "coverage": "Coverage",
        "use_in_model": "Use in Model",
        "confidence_level": "Confidence",
        "limitations": "Limitations",
    }

    return output.rename(columns=rename_map)


def prepare_correlation_table(df: pd.DataFrame) -> pd.DataFrame:
    """
    Prepare a compact correlation table.
    """

    if df is None or df.empty:
        return pd.DataFrame()

    corr = df.copy()

    if corr.columns[0].lower().startswith("unnamed"):
        corr = corr.rename(columns={corr.columns[0]: "variable"})

    if "variable" not in corr.columns:
        corr = corr.reset_index().rename(columns={"index": "variable"})

    wanted_rows = [
        "corruption_risk_scaled",
        "rule_of_law_risk_scaled",
        "political_instability_risk_scaled",
        "low_income_risk_scaled",
        "inflation_risk_scaled",
        "unemployment_risk_scaled",
        "tf_vulnerability_score",
        "fatf_grey_list",
        "predicted_grey_list_probability",
    ]

    wanted_cols = [
        "variable",
        "tf_vulnerability_score",
        "fatf_grey_list",
        "predicted_grey_list_probability",
    ]

    available_cols = [column for column in wanted_cols if column in corr.columns]
    corr = corr[corr["variable"].isin(wanted_rows)][available_cols]

    corr["Variable"] = corr["variable"].map(RISK_FEATURE_LABELS).fillna(
        corr["variable"]
    )

    rename_cols = {
        "tf_vulnerability_score": "Risk Score",
        "fatf_grey_list": "Grey List",
        "predicted_grey_list_probability": "Pred. Prob.",
    }

    corr = corr.rename(columns=rename_cols)

    for column in ["Risk Score", "Grey List", "Pred. Prob."]:
        if column in corr.columns:
            corr[column] = corr[column].apply(lambda value: safe_number(value, 3))

    return corr


# ---------------------------------------------------
# Report sections
# ---------------------------------------------------

def add_executive_summary(doc: Document, rankings_df: pd.DataFrame | None) -> None:
    """
    Add executive summary.
    """

    doc.add_heading("Executive Summary", level=1)

    top_country = "N/A"
    top_score = "N/A"

    if rankings_df is not None and not rankings_df.empty:
        top_row = rankings_df.sort_values(
            "tf_vulnerability_score",
            ascending=False,
        ).iloc[0]
        top_country = str(top_row["country"])
        top_score = safe_number(top_row["tf_vulnerability_score"], 2)

    add_paragraphs(
        doc,
        [
            (
                "This report presents a defensive open-source intelligence model for assessing "
                "country-level terrorism financing vulnerability. The model does not examine "
                "operational financing methods. Instead, it uses publicly available governance "
                "and economic indicators to create a structured risk score and compare that score "
                "against FATF grey-list status."
            ),
            (
                f"For the {STUDY_YEAR} target study year, the model ranked {top_country} as the "
                f"highest vulnerability country in the dataset, with a score of {top_score} out of 100."
            ),
            (
                "The model uses the latest available World Bank observations up to the study year "
                "when exact country-year values are unavailable. Indicator-year fields are preserved "
                "in the output tables to make this transparent."
            ),
            (
                "The study uses the most recent country-year data that can be supported consistently "
                "by the public source datasets. A 2025 or 2026 primary model year was not selected "
                "because the core World Bank governance and economic indicator datasets used in the "
                "quantitative model do not yet provide a complete and comparable 2025 or 2026 "
                "country-year panel. More recent official sources are therefore used for context "
                "and comparison, not as replacement model-year observations."
            ),
            (
                "The expanded project includes supplemental UCDP conflict-exposure data, ACLED recent "
                "conflict and unrest signals, Global Organized Crime Index enrichment, GDELT 2.0 batch public-news "
                "signal enrichment, top risk-driver interpretation, regional summaries, sanctions-exposure "
                "templates, FATF mutual-evaluation templates, official AML/CFT policy context, virtual-assets "
                "and Travel Rule context, threat finance intelligence-gap scoring, collection-priority scoring, "
                "and additional visual outputs. These additions improve interpretation, collection planning, "
                "and future extensibility, but they do not replace the primary country-level vulnerability score."
            ),
            (
                "The model should be interpreted as a screening and prioritization tool, not as a "
                "definitive determination of terrorism financing activity. Its value is in identifying "
                "countries where weak governance, political instability, conflict exposure, and economic "
                "stress may overlap with higher AML/CFT vulnerability."
            ),
        ],
    )


def add_project_overview(doc: Document) -> None:
    """
    Add project overview, research question, and source table.
    """

    doc.add_heading("Project Overview", level=1)

    add_paragraphs(
        doc,
        [
            (
                "This project develops a country-level quantitative OSINT risk model. The model is "
                "designed for defensive intelligence, financial-crime risk analysis, and academic "
                "research. It uses open-source data to estimate structural vulnerability rather than "
                "attempting to identify transactions, networks, or operational financing methods."
            ),
        ],
    )

    doc.add_heading("Research Question", level=2)
    doc.add_paragraph(
        "Can publicly available governance and economic indicators be used to estimate country-level "
        "terrorism financing vulnerability and partially explain FATF grey-list status?"
    )

    doc.add_heading("Data Sources", level=2)

    doc.add_paragraph(
        "The model uses World Bank governance and economic indicators, combined with FATF grey-list "
        "status as a validation target. Governance indicators are drawn from the World Bank Worldwide "
        "Governance Indicators framework, including the WGI 2025 Revision / WGI 2.0 methodological "
        "context where applicable. Economic indicators are drawn from the World Bank World Development "
        "Indicators."
    )

    doc.add_paragraph(
        f"The study year is set to {STUDY_YEAR} because it is the most recent year that can be supported "
        "by the project's core public indicator framework. Current FATF, Treasury, FATF, FIU, RUSI, GAO, "
        "CRS, EBA, UAE, Australia, and other official or policy sources are used as supplemental context."
    )

    source_rows = [
        [
            "World Bank WGI",
            "Control of Corruption",
            "control_of_corruption",
            "Governance integrity and corruption vulnerability.",
        ],
        [
            "World Bank WGI",
            "Rule of Law",
            "rule_of_law",
            "Legal and institutional enforcement strength.",
        ],
        [
            "World Bank WGI",
            "Political Stability",
            "political_stability",
            "Political instability and violence risk.",
        ],
        [
            "World Bank WDI",
            "GDP per Capita",
            "gdp_per_capita",
            "Economic development and income vulnerability.",
        ],
        [
            "World Bank WDI",
            "Inflation",
            "inflation",
            "Macroeconomic stress.",
        ],
        [
            "World Bank WDI",
            "Unemployment",
            "unemployment",
            "Labor-market stress.",
        ],
        [
            "FATF",
            "Grey-list status",
            "fatf_grey_list",
            "Validation target for model comparison.",
        ],
        [
            "UCDP",
            "Organized violence country-year data",
            "conflict_deaths_total",
            "Supplemental conflict-exposure context.",
        ],
        [
            "ACLED",
            "Recent conflict and political-violence event data",
            "acled_conflict_signal_score",
            "Supplemental recent conflict and unrest signal context.",
        ],
        [
            "Global Organized Crime Index",
            "Criminality and resilience scores",
            "organized_crime_score",
            "Supplemental organized-crime convergence context.",
        ],
        [
            "GDELT",
            "GDELT 2.0 batch public-news reporting signal",
            "gdelt_news_signal_score",
            "Supplemental public-news signal context.",
        ],
        [
            "World Bank country metadata",
            "Region and income metadata",
            "world_bank_region",
            "Regional grouping and summary analysis.",
        ],
        [
            "Manual templates",
            "Sanctions and FATF mutual-evaluation inputs",
            "supplemental templates",
            "Future structured AML/CFT enhancement layers.",
        ],
    ]

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"

    headers = ["Source", "Indicator", "Variable", "Purpose"]

    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header

    for row in source_rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = clean_table_value(value)

    format_word_table(table)
    doc.add_paragraph()

    doc.add_heading("Model Inputs and Analytical Role", level=2)

    add_paragraphs(
        doc,
        [
            (
                "The model separates inputs into four categories. First, core quantitative inputs are "
                "used directly in the baseline score: governance indicators, political stability, GDP per "
                "capita, inflation, and unemployment. Second, FATF grey-list status is used as a comparison "
                "and validation target, not as a direct scoring input. Third, UCDP conflict exposure, ACLED "
                "recent conflict and unrest signals, Global Organized Crime Index context, GDELT 2.0 batch public-news "
                "signal context, regional summaries, current FATF monitoring status, sanctions templates, "
                "mutual-evaluation templates, and policy sources are supplemental interpretation layers. "
                "Fourth, the threat finance intelligence layer converts baseline risk, live signals, source coverage, source staleness, and intelligence gaps "
                "into collection-priority outputs for analyst review."
            ),
            (
                "This separation is important for model governance. The baseline score remains transparent "
                "and reproducible because it is calculated from a small set of public country-level indicators. "
                "Supplemental layers help analysts interpret why a country may deserve additional review, while "
                "the intelligence layer helps structure what public-source information should be refreshed or "
                "validated next. These layers do not change the baseline score unless they are explicitly added "
                "in a future model version."
            ),
        ],
    )


def add_methodology(doc: Document, rankings_df: pd.DataFrame | None) -> None:
    """
    Add methodology and indicator-year coverage.
    """

    doc.add_heading("Methodology", level=1)

    add_paragraphs(
        doc,
        [
            (
                "The project uses a country-level risk scoring approach. The model pulls World Bank "
                "indicators through the public API, cleans the country-level dataset, removes regional "
                "aggregates, creates risk-oriented features, normalizes each feature to a common 0-to-1 "
                "scale, and applies weighted scoring. When exact study-year observations are unavailable "
                "for a country or indicator, the pipeline keeps the latest available observation up to the "
                "study year and records the year used."
            ),
            (
                "This latest-available rule is a deliberate data-quality choice. The study does not advance "
                "the primary model year to 2025 or 2026 because doing so would require mixing unavailable, "
                "incomplete, or non-comparable country-year observations across the core indicators."
            ),
            (
                "UCDP conflict exposure, ACLED recent conflict and unrest signals, Global Organized Crime "
                "Index enrichment, GDELT 2.0 batch public-news signal enrichment, regional summaries, sanctions exposure "
                "templates, FATF mutual-evaluation templates, official policy sources, and virtual-asset "
                "regulatory context are added as "
                "supplemental layers. The threat finance intelligence layer then uses baseline risk, live "
                "signals, source coverage, source staleness, and intelligence gaps to support collection-priority "
                "analysis. These layers improve interpretation and future extensibility, but they do not "
                "redefine the baseline terrorism-financing vulnerability score."
            ),
        ],
    )

    doc.add_heading("Process Summary", level=2)

    process_steps = [
        "Collect World Bank governance and economic data for all available countries using public API endpoints.",
        "Select the latest available observation up to the target study year for each indicator and preserve the source year.",
        "Remove regional aggregates and retain modeled country-level observations with complete core risk features.",
        "Convert raw indicators into risk-oriented variables so higher values consistently represent higher vulnerability.",
        "Normalize all risk features to a common 0-to-1 scale and apply transparent baseline weights.",
        "Rank countries by the final 0-to-100 terrorism-financing vulnerability score.",
        "Fit a logistic classification model to compare the risk features against FATF grey-list status.",
        "Add supplemental interpretation layers including risk tiers, sensitivity analysis, UCDP conflict exposure, ACLED recent-event signals, Global Organized Crime Index enrichment, GDELT 2.0 batch public-news signal enrichment, regional summaries, and current FATF comparison.",
        "Run threat finance intelligence-gap and collection-priority analysis to identify public-source review priorities.",
        "Generate model charts, supplemental charts, Word/PDF reports, and dashboard-ready outputs.",
    ]

    for step in process_steps:
        doc.add_paragraph(step, style="List Bullet")

    doc.add_heading("Risk Feature Construction", level=2)

    feature_rows = [
        [
            "Control of Corruption",
            "Lower control of corruption suggests weaker public-sector integrity.",
            "Inverted so higher values represent higher risk.",
        ],
        [
            "Rule of Law",
            "Lower rule of law suggests weaker legal and institutional enforcement.",
            "Inverted so higher values represent higher risk.",
        ],
        [
            "Political Stability",
            "Lower political stability suggests instability and violence risk.",
            "Inverted so higher values represent higher risk.",
        ],
        [
            "GDP per Capita",
            "Lower income may indicate greater economic vulnerability.",
            "Log transformed and inverted.",
        ],
        [
            "Inflation",
            "Higher inflation may indicate macroeconomic stress.",
            "Higher values treated as higher risk.",
        ],
        [
            "Unemployment",
            "Higher unemployment may indicate economic stress.",
            "Higher values treated as higher risk.",
        ],
    ]

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"

    for idx, header in enumerate(
        ["Indicator", "Risk Interpretation", "Model Treatment"]
    ):
        table.rows[0].cells[idx].text = header

    for row in feature_rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = clean_table_value(value)

    format_word_table(table)

    doc.add_heading("Risk Score Weights", level=2)

    for item in [
        "Corruption risk: 25%",
        "Rule of law risk: 20%",
        "Political instability risk: 25%",
        "Low income risk: 10%",
        "Inflation risk: 10%",
        "Unemployment risk: 10%",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Indicator Year Coverage", level=2)

    doc.add_paragraph(
        "The table below summarizes the year coverage retained in the final rankings file. This supports "
        "transparency when the latest available observation up to the study year is used."
    )

    if rankings_df is not None and not rankings_df.empty:
        year_rows = []

        for column in INDICATOR_YEAR_COLUMNS:
            if column not in rankings_df.columns:
                continue

            most_common_year = rankings_df[column].mode().iloc[0]
            exact_count = int((rankings_df[column] == STUDY_YEAR).sum())
            country_count = int(rankings_df[column].notna().sum())

            year_rows.append(
                {
                    "Indicator Year Field": column,
                    "Most Common Year": safe_year(most_common_year),
                    "Exact Study-Year Rows": safe_int(exact_count),
                    "Countries": safe_int(country_count),
                }
            )

        if year_rows:
            year_df = pd.DataFrame(year_rows)
            add_dataframe_table(
                doc,
                year_df,
                [
                    "Indicator Year Field",
                    "Most Common Year",
                    "Exact Study-Year Rows",
                    "Countries",
                ],
                [
                    "Indicator Year Field",
                    "Most Common Year",
                    "Exact Study-Year Rows",
                    "Countries",
                ],
            )
        else:
            doc.add_paragraph("No indicator-year fields were found in the rankings file.")


def add_top_country_rankings(doc: Document, rankings_df: pd.DataFrame | None) -> None:
    """
    Add top country rankings.
    """

    doc.add_heading("Top Country Risk Rankings", level=1)

    doc.add_paragraph(
        "The table below shows the top 20 countries by terrorism-financing vulnerability score. "
        "The score is a structural vulnerability indicator, not evidence of terrorist-financing activity."
    )

    if rankings_df is None or rankings_df.empty:
        doc.add_paragraph("Risk ranking data was not available.")
        return

    top_df = prepare_top_rankings(rankings_df)

    add_dataframe_table(
        doc,
        top_df,
        ["Rank", "Country", "Score", "Grey List", "Pred. Prob."],
        ["Rank", "Country", "Score", "Grey List", "Pred. Prob."],
        max_rows=MAX_TOP_RISK_ROWS,
    )

    doc.add_heading("How to Read the Rankings", level=2)

    add_paragraphs(
        doc,
        [
            (
                "The risk rank orders countries from highest to lowest structural vulnerability based on the "
                "weighted score. A higher score means the country has a stronger combination of weak governance, "
                "political instability, lower income conditions, inflation stress, or unemployment stress under "
                "the model assumptions."
            ),
            (
                "The Grey List column shows whether the country is coded as a FATF grey-list jurisdiction in the "
                "model target variable. The predicted probability is the model's estimated probability of grey-list "
                "status based on the selected risk features. It is useful for screening and comparison, but it is "
                "not a finding of terrorist-financing activity."
            ),
        ],
    )


def add_fatf_grey_list_background(doc: Document) -> None:
    """
    Add background on the FATF grey list and how it is used in the model.
    """

    doc.add_heading("FATF Grey-List Background", level=1)

    add_paragraphs(
        doc,
        [
            (
                "The Financial Action Task Force, or FATF, is the global standard-setting body for anti-money "
                "laundering, countering the financing of terrorism, and countering proliferation financing. FATF "
                "publishes public lists of jurisdictions with strategic weaknesses in their AML/CFT/CPF regimes."
            ),
            (
                "The term 'grey list' commonly refers to FATF's list of jurisdictions under increased monitoring. "
                "These jurisdictions are working with FATF to address strategic deficiencies in their regimes for "
                "countering money laundering, terrorist financing, and proliferation financing. When a jurisdiction "
                "is placed under increased monitoring, it has committed to resolve identified deficiencies within "
                "agreed timeframes and is subject to additional monitoring."
            ),
            (
                "Grey-list status does not mean that terrorist financing is occurring in a jurisdiction, and it "
                "should not be interpreted as evidence of terrorist-financing activity. It is a policy and regulatory "
                "monitoring status that reflects strategic AML/CFT/CPF deficiencies and a formal commitment to remediate them."
            ),
            (
                "In this project, FATF grey-list status is used as a comparison and validation target. The model asks "
                "whether broad public indicators such as corruption control, rule of law, political stability, income, "
                "inflation, and unemployment help explain which countries appear under increased monitoring. FATF status "
                "is not used as a direct scoring input in the baseline vulnerability score."
            ),
            (
                "This distinction is important. A country may receive a high model score because it has elevated structural "
                "vulnerability but may not be grey-listed by FATF. Conversely, a country may be grey-listed because of "
                "specific legal, supervisory, enforcement, beneficial ownership, financial-intelligence, or institutional "
                "deficiencies that are not fully captured by broad country-level public indicators."
            ),
        ],
    )


def add_modeling_results(doc: Document, rankings_df: pd.DataFrame | None) -> None:
    """
    Add modeling results and practical use case.
    """

    doc.add_heading("Modeling Results", level=1)

    doc.add_paragraph(
        "A logistic regression model was used to evaluate whether the scaled risk indicators help explain "
        "FATF grey-list status. Because FATF grey-listed countries represent a small minority of the dataset, "
        "the classification model used balanced class weights."
    )

    doc.add_paragraph(
        "The model input matrix consists of the six scaled risk features used in the baseline score. The target "
        "variable is the binary FATF grey-list flag. The model output is a predicted probability that allows the "
        "ranked vulnerability score to be compared with observed FATF monitoring status. This comparison is useful "
        "for validation and prioritization, but it should not be interpreted as a standalone compliance conclusion."
    )

    doc.add_heading("Top Predicted Grey-List Probabilities", level=2)

    if rankings_df is not None and not rankings_df.empty:
        probability_df = prepare_top_probabilities(rankings_df)

        add_dataframe_table(
            doc,
            probability_df,
            ["Country", "Risk Score", "Actual Grey List", "Pred. Probability"],
            ["Country", "Risk Score", "Actual Grey List", "Pred. Probability"],
            max_rows=MAX_TOP_PROBABILITY_ROWS,
        )

    doc.add_heading("Interpretation", level=2)

    add_paragraphs(
        doc,
        [
            (
                "The model should be interpreted as moderately predictive rather than definitive. FATF "
                "grey-list status is influenced by detailed AML/CFT evaluations, legal reforms, institutional "
                "capacity, enforcement effectiveness, and policy commitments. These factors are only partially "
                "captured by broad governance and economic indicators."
            ),
            (
                "The strongest intelligence value of the model is its use as an early-stage screening tool. "
                "It can help analysts identify countries where governance weakness, political instability, and "
                "economic stress overlap, creating a higher-risk environment that may warrant closer review."
            ),
            (
                "This model aligns with official public-sector guidance emphasizing that terrorist and illicit "
                "financing risk is best understood through multiple indicators and the totality of circumstances."
            ),
        ],
    )

    doc.add_heading("Practical Use Case for AML/CFT Screening", level=2)

    doc.add_paragraph(
        "The model can support early-stage geographic risk screening for financial crime, AML/CFT, and risk "
        "analytics teams. It is not intended to make final risk decisions, but it can help prioritize countries "
        "for deeper qualitative review and support structured discussion of governance, instability, conflict, "
        "organized-crime convergence context, GDELT 2.0 batch public-news signals, "
        "and economic stress indicators."
    )

    for item in [
        "Country-risk screening and geographic AML/CFT review.",
        "Prioritization of countries for deeper qualitative analysis.",
        "Comparison of internal risk views against public data indicators.",
        "Support for risk committee discussion and model documentation.",
        "Supplemental context for FATF monitoring, UCDP conflict exposure, ACLED recent-event signals, organized-crime convergence context, GDELT 2.0 batch public-news signals, and regional risk analysis.",
    ]:
        doc.add_paragraph(item, style="List Bullet")


def add_enhanced_analysis(
    doc: Document,
    risk_tier_df: pd.DataFrame | None,
    performance_df: pd.DataFrame | None,
    false_positive_df: pd.DataFrame | None,
    false_negative_df: pd.DataFrame | None,
    sensitivity_overlap_df: pd.DataFrame | None,
) -> None:
    """
    Add enhanced analysis section.
    """

    doc.add_heading("Enhanced Analysis", level=1)

    doc.add_paragraph(
        "This section expands the model interpretation beyond the country ranking. It translates numerical "
        "scores into risk tiers, summarizes model classification performance, reviews false positives and false "
        "negatives, and tests whether the rankings remain stable under alternative weighting assumptions."
    )

    doc.add_heading("Risk Tier Summary", level=2)

    if risk_tier_df is not None and not risk_tier_df.empty:
        display_df = prepare_risk_tier_summary(risk_tier_df)

        add_dataframe_table(
            doc,
            display_df,
            ["Risk Tier", "Countries", "Avg. Score", "Avg. Probability", "FATF Count"],
            ["Risk Tier", "Countries", "Avg. Score", "Avg. Probability", "FATF Count"],
        )

    doc.add_paragraph(
        "The tier results make the numerical score easier to communicate. Countries in the Very High and "
        "High tiers have the strongest overlap of governance weakness, political instability, and economic stress."
    )

    doc.add_heading("Model Performance Summary", level=2)

    if performance_df is not None and not performance_df.empty:
        display_df = prepare_model_performance(performance_df)

        add_dataframe_table(
            doc,
            display_df,
            ["Metric", "Value"],
            ["Metric", "Value"],
        )

    doc.add_paragraph(
        "The performance metrics show screening value but limited classification certainty. This is expected "
        "because FATF grey-list status is influenced by policy, supervisory, enforcement, and institutional "
        "factors not fully captured by broad country-level open-source indicators."
    )

    doc.add_heading("False Positive and False Negative Review", level=2)

    doc.add_paragraph(
        "The classification review uses a 50% predicted-probability threshold. False positives are not evidence "
        "of terrorist financing activity. They are countries where the model detects elevated structural "
        "vulnerability even though the country is not coded as FATF grey-listed. False negatives are FATF "
        "grey-listed countries that the model assigns a probability below the threshold."
    )

    if false_positive_df is not None:
        doc.add_paragraph(
            f"False positives identified: {len(false_positive_df)} countries. "
            "The table shows the highest-probability false positives."
        )

        display_df = prepare_false_output(false_positive_df)

        add_dataframe_table(
            doc,
            display_df,
            ["Country", "Rank", "Score", "Pred. Prob."],
            ["Country", "Rank", "Score", "Pred. Prob."],
            max_rows=MAX_FALSE_POSITIVE_ROWS,
        )

    if false_negative_df is not None:
        doc.add_paragraph(
            f"False negatives identified: {len(false_negative_df)} countries. "
            "The table shows FATF grey-listed countries below the 50% model threshold."
        )

        display_df = prepare_false_output(false_negative_df)

        add_dataframe_table(
            doc,
            display_df,
            ["Country", "Rank", "Score", "Pred. Prob."],
            ["Country", "Rank", "Score", "Pred. Prob."],
            max_rows=MAX_FALSE_NEGATIVE_ROWS,
        )

    doc.add_paragraph(
        "Analytically, false positives can be treated as candidates for further qualitative review, while false "
        "negatives show where FATF monitoring may reflect detailed institutional, legal, or supervisory weaknesses "
        "not fully represented in the selected public indicators."
    )

    doc.add_heading("Sensitivity Analysis", level=2)

    if sensitivity_overlap_df is not None and not sensitivity_overlap_df.empty:
        display_df = prepare_sensitivity_overlap(sensitivity_overlap_df)

        add_dataframe_table(
            doc,
            display_df,
            ["Scenario", "Top-20 Overlap", "Baseline Count", "Overlap %"],
            ["Scenario", "Top-20 Overlap", "Baseline Count", "Overlap %"],
        )

    doc.add_paragraph(
        "The sensitivity results indicate whether the top country group remains reasonably stable across "
        "weighting assumptions. This strengthens the usefulness of the model as a transparent screening tool "
        "while still showing that exact rank order depends on how governance, instability, and economic stress "
        "are weighted."
    )


def add_top_risk_drivers(doc: Document, top_drivers_df: pd.DataFrame | None) -> None:
    """
    Add top risk drivers section.
    """

    doc.add_heading("Top Risk Drivers", level=1)

    doc.add_paragraph(
        "This section improves model interpretability by identifying the largest risk-score contributors for "
        "the highest-ranked countries. The driver labels explain which normalized risk factors most strongly "
        "contributed to each country's vulnerability score."
    )

    if top_drivers_df is not None and not top_drivers_df.empty:
        display_df = prepare_risk_drivers(top_drivers_df)

        add_dataframe_table(
            doc,
            display_df,
            ["Rank", "Country", "Score", "Driver 1", "Driver 2", "Driver 3"],
            ["Rank", "Country", "Score", "Driver 1", "Driver 2", "Driver 3"],
            max_rows=MAX_TOP_DRIVER_ROWS,
        )

    doc.add_paragraph(
        "Risk drivers should be read as model explanations, not as independent proof of terrorist-financing "
        "activity. They show why the scoring framework placed a country high in the ranking."
    )


def add_conflict_supplement(doc: Document, conflict_df: pd.DataFrame | None) -> None:
    """
    Add conflict event data supplement.
    """

    doc.add_heading("Conflict Event Data Supplement", level=1)

    doc.add_paragraph(
        "This supplemental section adds country-level organized-violence context using the UCDP country-year "
        "organized-violence dataset. Conflict exposure is relevant to terrorism-financing vulnerability because "
        "conflict, instability, and weak state capacity can affect AML/CFT risk environments. The conflict "
        "supplement does not replace the main vulnerability score and is not evidence of terrorist-financing activity."
    )

    if conflict_df is None or conflict_df.empty:
        doc.add_paragraph("Conflict supplement data was not available.")
        return

    matched_count = (
        int(conflict_df["ucdp_match_found"].sum())
        if "ucdp_match_found" in conflict_df.columns
        else int(conflict_df["ucdp_year"].notna().sum())
        if "ucdp_year" in conflict_df.columns
        else len(conflict_df)
    )

    if "organized_violence_active" in conflict_df.columns:
        active_count = int(conflict_df["organized_violence_active"].sum())
    elif "conflict_deaths_total" in conflict_df.columns:
        active_count = int((conflict_df["conflict_deaths_total"] > 0).sum())
    else:
        active_count = 0

    total_deaths = (
        int(conflict_df["conflict_deaths_total"].sum())
        if "conflict_deaths_total" in conflict_df.columns
        else 0
    )

    doc.add_paragraph(
        f"Within the modeled dataset, {matched_count:,} countries were matched to UCDP country-year records. "
        f"The supplement identifies {active_count:,} countries with recorded organized-violence deaths and "
        f"{total_deaths:,} total recorded organized-violence deaths in the matched supplement."
    )

    doc.add_heading("Top Countries by Recorded Organized-Violence Deaths", level=2)

    display_df = prepare_conflict_table(conflict_df)

    add_dataframe_table(
        doc,
        display_df,
        ["Country", "Rank", "Risk Score", "UCDP Year", "Deaths", "Conflict Events", "Category"],
        ["Country", "Rank", "Risk Score", "UCDP Year", "Deaths", "Conflict Events", "Category"],
        max_rows=MAX_CONFLICT_ROWS,
    )

    doc.add_paragraph(
        "Large conflict-death counts are treated as supplemental instability context. They should be combined "
        "with governance, AML/CFT, sanctions, FATF, and qualitative context before any country-level risk "
        "conclusion is drawn."
    )


def add_additional_risk_factors(
    doc: Document,
    regional_df: pd.DataFrame | None,
    summary_df: pd.DataFrame | None,
    sanctions_df: pd.DataFrame | None,
    mutual_eval_df: pd.DataFrame | None,
) -> None:
    """
    Add additional risk factor analysis.
    """

    doc.add_heading("Additional Risk Factor Analysis", level=1)

    doc.add_paragraph(
        "This section summarizes supplemental risk-factor outputs added to make the project more extensible. "
        "Regional summaries provide geographic structure, while the sanctions and FATF mutual-evaluation "
        "templates create input files for future structured enhancements. ACLED, Global Organized Crime "
        "Index, and GDELT outputs are handled as separate supplemental signal and convergence layers and "
        "are summarized in dedicated sections below. These outputs are supplemental and do not replace "
        "the baseline model score."
    )

    doc.add_heading("Regional Risk Summary", level=2)

    if regional_df is not None and not regional_df.empty:
        display_df = regional_df.copy()

        rename_map = {
            "world_bank_region": "Region",
            "region": "Region",
            "country_count": "Country Count",
            "average_risk_score": "Average Risk Score",
            "median_risk_score": "Median Risk Score",
            "fatf_grey_list_count": "FATF Grey List Count",
            "current_fatf_monitoring_count": "Current FATF Monitoring Count",
            "highest_risk_country": "Highest Risk Country",
        }

        display_df = display_df.rename(columns=rename_map)

        for column in ["Average Risk Score", "Median Risk Score"]:
            if column in display_df.columns:
                display_df[column] = display_df[column].apply(
                    lambda value: safe_number(value, 2)
                )

        add_dataframe_table(
            doc,
            display_df,
            [
                "Region",
                "Country Count",
                "Average Risk Score",
                "Median Risk Score",
                "FATF Grey List Count",
                "Current FATF Monitoring Count",
                "Highest Risk Country",
            ],
            [
                "Region",
                "Country Count",
                "Average Risk Score",
                "Median Risk Score",
                "FATF Grey List Count",
                "Current FATF Monitoring Count",
                "Highest Risk Country",
            ],
        )

    doc.add_heading("Sanctions and FATF Mutual Evaluation Templates", level=2)

    sanctions_rows = 0 if sanctions_df is None else len(sanctions_df)
    mutual_eval_rows = 0 if mutual_eval_df is None else len(mutual_eval_df)

    doc.add_paragraph(
        f"Sanctions exposure supplement created with {sanctions_rows:,} country rows.",
        style="List Bullet",
    )
    doc.add_paragraph(
        f"FATF mutual-evaluation supplement created with {mutual_eval_rows:,} country rows.",
        style="List Bullet",
    )

    doc.add_paragraph(
        "These files are intentionally created as country-level input templates. The 173 rows indicate that every "
        "modeled country has a placeholder record that can be populated later with structured sanctions, FATF "
        "technical-compliance, or FATF effectiveness data. The current template rows should not be interpreted as "
        "evidence of sanctions exposure, FATF deficiencies, or terrorism-financing activity."
    )

    doc.add_paragraph(
        "This design keeps the baseline model transparent. The current vulnerability score continues to rely only "
        "on the core governance and economic indicators. Sanctions exposure and FATF mutual-evaluation data can be "
        "added later as separate supplemental layers once reliable, country-comparable source values are collected."
    )

    doc.add_heading("Additional Risk Factor Output Summary", level=2)

    if summary_df is not None and not summary_df.empty:
        display_df = summary_df.rename(
            columns={
                "factor": "Factor",
                "status": "Status",
                "records_or_groups": "Records Or Groups",
                "output_file": "Output File",
            }
        ).copy()

        supplemental_rows = []

        if ORGANIZED_CRIME_ENRICHMENT_FILE.exists():
            organized_crime_df = read_csv_if_exists(ORGANIZED_CRIME_ENRICHMENT_FILE)
            organized_crime_records = (
                len(organized_crime_df)
                if organized_crime_df is not None
                else 0
            )

            if not display_df["Factor"].astype(str).str.contains(
                "organized crime",
                case=False,
                na=False,
            ).any():
                supplemental_rows.append(
                    {
                        "Factor": "Global Organized Crime Index enrichment",
                        "Status": "Created if organized-crime enrichment output exists",
                        "Records Or Groups": organized_crime_records,
                        "Output File": str(ORGANIZED_CRIME_ENRICHMENT_FILE),
                    }
                )

        if GDELT_NEWS_SIGNAL_ENRICHMENT_FILE.exists():
            gdelt_df = read_csv_if_exists(GDELT_NEWS_SIGNAL_ENRICHMENT_FILE)
            gdelt_records = len(gdelt_df) if gdelt_df is not None else 0

            if not display_df["Factor"].astype(str).str.contains(
                "gdelt",
                case=False,
                na=False,
            ).any():
                supplemental_rows.append(
                    {
                        "Factor": "GDELT 2.0 batch public-news signal",
                        "Status": "Created if GDELT batch enrichment output exists",
                        "Records Or Groups": gdelt_records,
                        "Output File": str(GDELT_NEWS_SIGNAL_ENRICHMENT_FILE),
                    }
                )

        if supplemental_rows:
            display_df = pd.concat(
                [display_df, pd.DataFrame(supplemental_rows)],
                ignore_index=True,
            )

        if "Output File" in display_df.columns:
            display_df["Output File"] = display_df["Output File"].apply(
                shorten_output_path
            )

        add_dataframe_table(
            doc,
            display_df,
            ["Factor", "Status", "Records Or Groups", "Output File"],
            ["Factor", "Status", "Records", "Output"],
        )


def add_threat_finance_intelligence_gap_analysis(
    doc: Document,
    collection_priority_df: pd.DataFrame | None,
    source_register_df: pd.DataFrame | None,
) -> None:
    """
    Add threat finance intelligence gap and collection prioritization section.
    """

    doc.add_heading(
        "Threat Finance Intelligence Gap and Collection Prioritization",
        level=1,
    )

    add_paragraphs(
        doc,
        [
            (
                "This section adds an intelligence-cycle layer to the terrorism-financing "
                "OSINT model. The baseline model estimates country-level structural "
                "vulnerability. The threat finance intelligence layer asks a different "
                "question: what information should an analyst collect next to validate, "
                "challenge, or update the assessment?"
            ),
            (
                "The collection-priority score combines baseline terrorism-financing risk, "
                "recent or live signal strength, source staleness, intelligence gaps, and "
                "decision relevance. The result is not a terrorism determination, sanctions "
                "finding, or allegation of illicit activity. It is a public-source collection "
                "planning aid for analyst review."
            ),
            (
                "This section is designed to move the project from a static country-risk "
                "ranking toward a more mature intelligence decision-support workflow: "
                "risk score, intelligence gap, collection priority, and analyst review note."
            ),
        ],
    )

    if collection_priority_df is None or collection_priority_df.empty:
        doc.add_paragraph(
            "Threat finance collection-priority output was not available."
        )
        return

    priority_count = len(collection_priority_df)

    immediate_count = (
        int(
            (
                collection_priority_df["collection_priority_level"]
                == "Immediate Review"
            ).sum()
        )
        if "collection_priority_level" in collection_priority_df.columns
        else 0
    )

    high_count = (
        int(
            (
                collection_priority_df["collection_priority_level"]
                == "High Priority"
            ).sum()
        )
        if "collection_priority_level" in collection_priority_df.columns
        else 0
    )

    avg_gap = (
        collection_priority_df["intelligence_gap_score"].mean()
        if "intelligence_gap_score" in collection_priority_df.columns
        else None
    )

    avg_priority = (
        collection_priority_df["collection_priority_score"].mean()
        if "collection_priority_score" in collection_priority_df.columns
        else None
    )

    doc.add_paragraph(
        f"The collection-priority output contains {priority_count:,} country records. "
        f"{immediate_count:,} countries are classified as Immediate Review and "
        f"{high_count:,} countries are classified as High Priority. "
        f"The average intelligence-gap score is {safe_number(avg_gap, 2)}, and "
        f"the average collection-priority score is {safe_number(avg_priority, 2)}."
    )

    doc.add_heading("Top Collection Priorities", level=2)

    display_df = prepare_collection_priority_table(collection_priority_df)

    add_dataframe_table(
        doc,
        display_df,
        [
            "Country",
            "ISO3",
            "Baseline Risk",
            "Risk Level",
            "Live Signal",
            "Intel Gap",
            "Gap Level",
            "Collection Priority",
            "Priority Level",
        ],
        [
            "Country",
            "ISO3",
            "Baseline Risk",
            "Risk Level",
            "Live Signal",
            "Intel Gap",
            "Gap Level",
            "Priority",
            "Priority Level",
        ],
        max_rows=MAX_COLLECTION_PRIORITY_ROWS,
    )

    doc.add_paragraph(
        "Detailed recommended collection-focus text is preserved in the CSV output "
        "and analyst-note field. It is omitted from the PDF table to keep the "
        "collection-priority table readable."
    )

    doc.add_heading("Analyst Review Interpretation", level=2)

    add_paragraphs(
        doc,
        [
            (
                "Countries at the top of the collection-priority ranking are not being "
                "identified as terrorism-financing actors or confirmed threat locations. "
                "They are countries where the model indicates that structural risk, source "
                "staleness, live-signal assumptions, and intelligence gaps justify deeper "
                "public-source review."
            ),
            (
                "The recommended collection focus should be treated as a tasking guide. "
                "For example, a country may require refreshed sanctions review, updated "
                "conflict reporting, AML/CFT effectiveness checks, enforcement-action review, "
                "or additional terrorism-financing reporting validation."
            ),
            (
                "The analyst note generated by the model is intentionally conservative. "
                "It explains the review priority while stating that the output does not "
                "constitute a determination of terrorism, criminality, sanctions exposure, "
                "or illicit activity."
            ),
        ],
    )

    doc.add_heading("Source Register", level=2)

    if source_register_df is not None and not source_register_df.empty:
        source_display_df = prepare_source_register_table(source_register_df)

        add_dataframe_table(
            doc,
            source_display_df,
            [
                "Source",
                "Type",
                "Coverage",
                "Use in Model",
                "Confidence",
                "Limitations",
            ],
            [
                "Source",
                "Type",
                "Coverage",
                "Use in Model",
                "Confidence",
                "Limitations",
            ],
            max_rows=MAX_SOURCE_REGISTER_ROWS,
        )

        doc.add_paragraph(
            "The source register is included to support model transparency. It documents "
            "the public-source categories used or planned in the intelligence layer, their "
            "analytical role, confidence level, and limitations."
        )
    else:
        doc.add_paragraph("Source register output was not available.")


def add_acled_signal_supplement(
    doc: Document,
    acled_df: pd.DataFrame | None,
) -> None:
    """
    Add ACLED recent conflict and unrest signal supplement.
    """

    doc.add_heading("ACLED Recent Conflict and Unrest Signal Supplement", level=1)

    add_paragraphs(
        doc,
        [
            (
                "This supplemental section summarizes ACLED recent conflict, unrest, "
                "and political-violence signals. ACLED is used as a near-real-time "
                "early-warning context layer for the live-signal and collection-priority "
                "workflow. It does not replace the baseline terrorism-financing "
                "vulnerability score."
            ),
            (
                "The ACLED signal should not be interpreted as evidence of terrorism "
                "financing, criminality, sanctions exposure, or illicit activity. It is "
                "a public-source context signal that can help analysts decide whether "
                "recent instability conditions warrant refreshed review."
            ),
        ],
    )

    if acled_df is None or acled_df.empty:
        doc.add_paragraph("ACLED signal enrichment output was not available.")
        return

    country_count = len(acled_df)

    total_events = (
        int(pd.to_numeric(acled_df["acled_event_count"], errors="coerce").fillna(0).sum())
        if "acled_event_count" in acled_df.columns
        else 0
    )

    total_fatalities = (
        int(pd.to_numeric(acled_df["acled_fatalities"], errors="coerce").fillna(0).sum())
        if "acled_fatalities" in acled_df.columns
        else 0
    )

    high_or_severe_count = 0
    if "acled_signal_level" in acled_df.columns:
        high_or_severe_count = int(
            acled_df["acled_signal_level"].isin(["High", "Severe"]).sum()
        )

    doc.add_paragraph(
        f"The ACLED enrichment output contains {country_count:,} country or territory "
        f"records, {total_events:,} recent event records after aggregation, and "
        f"{total_fatalities:,} recorded fatalities in the configured lookback window. "
        f"{high_or_severe_count:,} records are classified as High or Severe in the "
        f"ACLED signal layer."
    )

    doc.add_heading("Top ACLED Recent Conflict and Unrest Signals", level=2)

    display_df = prepare_acled_signal_table(acled_df)

    add_dataframe_table(
        doc,
        display_df,
        [
            "Country",
            "ISO / Code",
            "Events",
            "Fatalities",
            "Political Violence",
            "Unrest",
            "ACLED Signal",
            "Signal Level",
        ],
        [
            "Country",
            "ISO / Code",
            "Events",
            "Fatalities",
            "Political Violence",
            "Unrest",
            "ACLED Signal",
            "Signal Level",
        ],
        max_rows=MAX_ACLED_SIGNAL_ROWS,
    )

    doc.add_paragraph(
        "Countries or territories appearing in this table should be treated as "
        "recent instability review candidates only. They are not being identified "
        "as terrorism-financing jurisdictions or threat actors."
    )


def add_organized_crime_supplement(
    doc: Document,
    organized_crime_df: pd.DataFrame | None,
) -> None:
    """
    Add Global Organized Crime Index supplemental context section.
    """

    doc.add_heading("Global Organized Crime Index Supplement", level=1)

    add_paragraphs(
        doc,
        [
            (
                "This supplemental section adds country-level organized-crime exposure "
                "and resilience context using the Global Organized Crime Index. The "
                "organized-crime layer is used as convergence context for the "
                "threat-finance intelligence and collection-priority workflow."
            ),
            (
                "The organized-crime score does not replace the baseline terrorism-"
                "financing vulnerability score. It should not be interpreted as evidence "
                "of terrorism financing, criminality, sanctions exposure, illicit finance, "
                "or activity by any specific person, entity, account, wallet, or network."
            ),
        ],
    )

    if organized_crime_df is None or organized_crime_df.empty:
        doc.add_paragraph("Organized-crime enrichment output was not available.")
        return

    record_count = len(organized_crime_df)

    matched_count = (
        int(
            pd.to_numeric(
                organized_crime_df["organized_crime_source_present"],
                errors="coerce",
            )
            .fillna(0)
            .sum()
        )
        if "organized_crime_source_present" in organized_crime_df.columns
        else int(organized_crime_df["organized_crime_score"].notna().sum())
        if "organized_crime_score" in organized_crime_df.columns
        else 0
    )

    avg_score = (
        organized_crime_df["organized_crime_score"].mean()
        if "organized_crime_score" in organized_crime_df.columns
        else None
    )

    high_or_above_count = 0
    if "organized_crime_level" in organized_crime_df.columns:
        high_or_above_count = int(
            organized_crime_df["organized_crime_level"].isin(
                ["High", "Very High"]
            ).sum()
        )

    latest_year = "N/A"
    if (
        "oc_index_year" in organized_crime_df.columns
        and organized_crime_df["oc_index_year"].notna().any()
    ):
        latest_year = safe_year(organized_crime_df["oc_index_year"].max())

    doc.add_paragraph(
        f"The organized-crime enrichment output contains {record_count:,} modeled "
        f"country records, with {matched_count:,} records matched to the Global "
        f"Organized Crime Index. The latest index year represented is {latest_year}. "
        f"The average organized-crime vulnerability score is {safe_number(avg_score, 2)}, "
        f"and {high_or_above_count:,} records are classified as High or Very High."
    )

    doc.add_heading("Top Organized-Crime Vulnerability Scores", level=2)

    display_df = prepare_organized_crime_table(organized_crime_df)

    add_dataframe_table(
        doc,
        display_df,
        [
            "Country",
            "ISO3",
            "Index Year",
            "Criminality",
            "Resilience",
            "Exposure",
            "Resilience Gap",
            "OC Vulnerability",
            "OC Level",
        ],
        [
            "Country",
            "ISO3",
            "Index Year",
            "Criminality",
            "Resilience",
            "Exposure",
            "Resilience Gap",
            "OC Score",
            "OC Level",
        ],
        max_rows=MAX_ORGANIZED_CRIME_ROWS,
    )

    doc.add_paragraph(
        "Countries appearing in this table should be treated as organized-crime "
        "context review candidates only. The table is useful for understanding "
        "where illicit-market exposure and resilience gaps may overlap with "
        "structural vulnerability, but it is not a terrorism-financing or "
        "criminality determination."
    )


def add_gdelt_news_signal_supplement(
    doc: Document,
    gdelt_df: pd.DataFrame | None,
) -> None:
    """
    Add GDELT 2.0 batch public-news signal supplemental context section.
    """

    doc.add_heading("GDELT 2.0 Batch Public-News Signal Supplement", level=1)

    add_paragraphs(
        doc,
        [
            (
                "This supplemental section adds country-level public-news signal context "
                "using locally filtered GDELT 2.0 Global Knowledge Graph batch files. "
                "The GDELT layer is used as media-reporting and open-source signal "
                "context for the threat-finance intelligence and collection-priority "
                "workflow."
            ),
            (
                "The GDELT signal does not replace the baseline terrorism-financing "
                "vulnerability score. It should not be interpreted as evidence of terrorism "
                "financing, criminality, sanctions exposure, illicit finance, or activity by "
                "any specific person, entity, account, wallet, or network. Public reporting "
                "can include duplication, media bias, translation artifacts, false positives, "
                "and uneven country coverage."
            ),
            (
                "This project uses GDELT 2.0 batch files rather than one country-by-country "
                "DOC API query per modeled jurisdiction. The batch approach reduces "
                "rate-limit exposure and supports repeatable local filtering."
            ),
        ],
    )

    if gdelt_df is None or gdelt_df.empty:
        doc.add_paragraph("GDELT 2.0 batch news signal enrichment output was not available.")
        return

    record_count = len(gdelt_df)

    matched_count = (
        int((pd.to_numeric(gdelt_df["gdelt_source_present"], errors="coerce").fillna(0) == 1).sum())
        if "gdelt_source_present" in gdelt_df.columns
        else int((pd.to_numeric(gdelt_df["gdelt_matching_record_count"], errors="coerce").fillna(0) > 0).sum())
        if "gdelt_matching_record_count" in gdelt_df.columns
        else 0
    )

    total_records = (
        pd.to_numeric(gdelt_df["gdelt_matching_record_count"], errors="coerce")
        .fillna(0)
        .sum()
        if "gdelt_matching_record_count" in gdelt_df.columns
        else pd.to_numeric(gdelt_df["gdelt_article_count"], errors="coerce")
        .fillna(0)
        .sum()
        if "gdelt_article_count" in gdelt_df.columns
        else 0
    )

    total_sources = (
        pd.to_numeric(gdelt_df["gdelt_matching_source_count"], errors="coerce")
        .fillna(0)
        .sum()
        if "gdelt_matching_source_count" in gdelt_df.columns
        else 0
    )

    batch_file_count = "N/A"
    if (
        "gdelt_batch_file_count" in gdelt_df.columns
        and gdelt_df["gdelt_batch_file_count"].notna().any()
    ):
        batch_file_count = safe_int(gdelt_df["gdelt_batch_file_count"].max())

    avg_signal = (
        gdelt_df["gdelt_news_signal_score"].mean()
        if "gdelt_news_signal_score" in gdelt_df.columns
        else None
    )

    high_or_above_count = 0
    if "gdelt_news_signal_level" in gdelt_df.columns:
        high_or_above_count = int(
            gdelt_df["gdelt_news_signal_level"].isin(
                ["High", "Very High"]
            ).sum()
        )

    doc.add_paragraph(
        f"The GDELT 2.0 batch enrichment output contains {record_count:,} modeled "
        f"country records. {matched_count:,} countries have at least one matched "
        f"GDELT signal record. The enrichment processed {batch_file_count} GDELT "
        f"batch files, identified {safe_number(total_records, 0)} matched GKG "
        f"country records, and counted {safe_number(total_sources, 0)} matched "
        f"source references. The average GDELT signal score is "
        f"{safe_number(avg_signal, 2)}, and {high_or_above_count:,} records are "
        f"classified as High or Very High."
    )

    doc.add_heading("Top GDELT 2.0 Batch Public-News Signal Scores", level=2)

    display_df = prepare_gdelt_signal_table(gdelt_df)

    add_dataframe_table(
        doc,
        display_df,
        [
            "Country",
            "ISO3",
            "Batch Files",
            "Matched Records",
            "Matched Sources",
            "Average Tone",
            "GDELT Signal",
            "Signal Level",
            "Batch Status",
        ],
        [
            "Country",
            "ISO3",
            "Files",
            "Records",
            "Sources",
            "Avg. Tone",
            "GDELT Signal",
            "Signal Level",
            "Status",
        ],
        max_rows=MAX_GDELT_SIGNAL_ROWS,
    )

    doc.add_paragraph(
        "Countries appearing in this table should be treated as public-news review "
        "candidates only. The table is useful for understanding where recent GDELT "
        "media-reporting signals may justify refreshed public-source review, but it "
        "is not a terrorism-financing, sanctions, or criminality determination."
    )


def add_current_data_supplement(
    doc: Document,
    current_df: pd.DataFrame | None,
    current_fatf_df: pd.DataFrame | None,
) -> None:
    """
    Add current data supplement.
    """

    doc.add_heading("Current Data Supplement and Forward-Looking Risk Context", level=1)

    doc.add_paragraph(
        "This section compares the model outputs with the current FATF increased-monitoring status configured "
        "in the project and current economic indicators collected as a supplement. The supplement does not "
        "replace the main model; it provides an additional current-status comparison layer."
    )

    doc.add_paragraph(
        "This distinction matters because FATF monitoring status and official risk guidance can be more current "
        "than the structured country-year indicator datasets used for model scoring. The project therefore separates "
        "the quantitative model year from forward-looking context."
    )

    doc.add_paragraph(
        "In this run, the current economic supplement uses the same target year as the baseline economic indicators. "
        "As a result, the economic change fields may show 0.00 when the latest available current value is the same "
        "observation used in the baseline model. The main analytical value of this section is therefore the current "
        "FATF status comparison and the ability to refresh this layer in future model runs."
    )

    if current_fatf_df is not None and not current_fatf_df.empty:
        current_count = (
            int(current_fatf_df["current_fatf_increased_monitoring"].sum())
            if "current_fatf_increased_monitoring" in current_fatf_df.columns
            else 0
        )

        difference_count = (
            int(current_fatf_df["fatf_status_changed"].sum())
            if "fatf_status_changed" in current_fatf_df.columns
            else 0
        )

        match_count = (
            int(current_fatf_df["model_signal_matches_current_fatf"].sum())
            if "model_signal_matches_current_fatf" in current_fatf_df.columns
            else 0
        )

        difference_noun = pluralize(difference_count, "country", "countries")

        doc.add_paragraph(
            f"Within the modeled dataset, {current_count:,} countries are coded as currently under FATF increased "
            f"monitoring in the project configuration. The current-status comparison flags {difference_count:,} "
            f"{difference_noun} where the model FATF target and current FATF status differ. The model's 50% probability "
            f"signal matches current FATF status for {match_count:,} countries out of {len(current_fatf_df):,}."
        )

        doc.add_heading("Current FATF Monitored Jurisdictions by Model Probability", level=2)

        display_df = prepare_current_fatf_table(current_fatf_df)

        add_dataframe_table(
            doc,
            display_df,
            ["Country", "Rank", "Score", "Pred. Prob.", "Model FATF", "Current FATF"],
            ["Country", "Rank", "Score", "Pred. Prob.", "Model FATF", "Current FATF"],
            max_rows=MAX_CURRENT_FATF_ROWS,
        )

    doc.add_heading("Economic Indicator Change Summary", level=2)

    if current_df is not None and not current_df.empty:
        change_rows = []

        change_specs = [
            ("gdp_per_capita_change_pct", "GDP per Capita Change %"),
            ("inflation_change_points", "Inflation Change, Points"),
            ("unemployment_change_points", "Unemployment Change, Points"),
        ]

        for column, label in change_specs:
            if column not in current_df.columns:
                continue

            change_rows.append(
                {
                    "Measure": label,
                    "Mean": safe_number(current_df[column].mean(), 2),
                    "Median": safe_number(current_df[column].median(), 2),
                    "Available Rows": safe_int(current_df[column].notna().sum()),
                }
            )

        if change_rows:
            display_df = pd.DataFrame(change_rows)

            add_dataframe_table(
                doc,
                display_df,
                ["Measure", "Mean", "Median", "Available Rows"],
                ["Measure", "Mean", "Median", "Available Rows"],
            )
        else:
            doc.add_paragraph("Current economic change columns were not available.")

    doc.add_paragraph(
        "These change measures are supplemental. In this run, the baseline model and current economic supplement "
        "rely on the same latest-available economic observations, so the reported changes are 0.00. In future runs, "
        "this table can help identify whether economic conditions have shifted relative to the model's main indicator "
        "values. These changes should not be interpreted as direct evidence of terrorist-financing activity."
    )


def add_official_designation_context(doc: Document) -> None:
    """
    Add official terrorism designation context section.
    """

    doc.add_heading("Official Terrorism Designation Context by Region", level=1)

    add_paragraphs(
        doc,
        [
            (
                "A regional designation context layer could add interpretive value to the project, but it should "
                "be framed carefully. The report should avoid describing a broad list of 'confirmed terrorist "
                "groups.' A more accurate description is official terrorism designation context, because designations "
                "depend on the legal authority, jurisdiction, date, and list used."
            ),
            (
                "This section is therefore treated as supplemental context only. Future versions could summarize "
                "the number of officially designated entities by region using official sources such as the UN "
                "Security Council ISIL/Da'esh and Al-Qaida sanctions list, the U.S. Department of State Foreign "
                "Terrorist Organizations list, and comparable official national or multilateral designation lists."
            ),
            (
                "The safest use in this model is a high-level regional count or source-reference table. Detailed "
                "group profiles, operational descriptions, tactical details, or financing-method descriptions are "
                "outside the scope of this defensive country-level screening model."
            ),
        ],
    )


def add_south_africa_fic_context(doc: Document) -> None:
    """
    Add South Africa FIC case study context.
    """

    doc.add_heading("South Africa Case Study: FATF Grey-List Remediation and FIU Context", level=1)

    add_paragraphs(
        doc,
        [
            (
                "South Africa provides a useful country-level example of how FATF grey-list status can reflect "
                "institutional and supervisory deficiencies rather than direct evidence of terrorist-financing "
                "activity. The Financial Intelligence Centre Annual Report 2024/25 states that South Africa was "
                "grey listed by FATF in February 2023 and that, by the end of the reporting year, the country had "
                "addressed 20 of 22 FATF action-plan items."
            ),
            (
                "The report also notes that FATF announced in June 2025 that South Africa had substantially "
                "completed all 22 action-plan items, subject to verification. This supports the study's treatment "
                "of FATF status as a policy and regulatory benchmark rather than a direct terrorism-financing "
                "activity measure."
            ),
            (
                "The South Africa case also shows why country-level structural indicators should be interpreted "
                "with institutional context. A jurisdiction's FATF status may change as legal reforms, supervisory "
                "improvements, beneficial ownership controls, enforcement effectiveness, and financial intelligence "
                "capabilities develop over time."
            ),
        ],
    )


def add_treasury_ntfra_context(doc: Document) -> None:
    """
    Add Treasury National Terrorist Financing Risk Assessment context.
    """

    doc.add_heading("U.S. Treasury 2024 National Terrorist Financing Risk Assessment Context", level=1)

    add_paragraphs(
        doc,
        [
            (
                "The U.S. Treasury 2024 National Terrorist Financing Risk Assessment provides official government "
                "context for terrorist-financing threats, vulnerabilities, financial channels, and risk-assessment "
                "methodology. It is useful for this study because it reinforces the need to interpret country-level "
                "vulnerability alongside more detailed financial-sector, enforcement, and typology information."
            ),
            (
                "The assessment supports the report's limitations because it discusses multiple channels and "
                "vulnerability areas that the baseline model does not directly measure, including money services "
                "businesses, person-to-person payments, unregistered money transmission, cash, banks, virtual assets, "
                "non-profit organizations, crowdfunding, and online fundraising."
            ),
            (
                "This context strengthens the study's conclusion that the OSINT model is a screening framework, "
                "not a complete terrorist-financing detection system. It can help prioritize geographic review, "
                "but higher-confidence assessment requires transaction-level indicators, institutional controls, "
                "law enforcement information, financial intelligence, and qualitative analysis."
            ),
        ],
    )


def add_treasury_strategy_context(doc: Document) -> None:
    """
    Add Treasury illicit finance strategy context.
    """

    doc.add_heading("U.S. Treasury 2024 Illicit Finance Strategy Context", level=1)

    add_paragraphs(
        doc,
        [
            (
                "The U.S. Treasury 2024 National Strategy for Combating Terrorist and Other Illicit Financing "
                "provides broader strategic context for this project. It describes the U.S. government's approach "
                "to denying illicit actors access to the funds and resources needed to carry out harmful activities "
                "or profit from crimes."
            ),
            (
                "This strategy fits the model's defensive framing because it emphasizes risk-based prioritization, "
                "AML/CFT modernization, transparency, supervision, enforcement, and disruption of illicit-finance "
                "networks. The model's value is not final classification; it is a structured way to prioritize "
                "geographic risk review using transparent public indicators."
            ),
            (
                "The strategy also supports future expansion into legal, supervisory, transparency, sanctions, "
                "and enforcement indicators. Those indicators should remain separate supplemental layers until "
                "reliable and country-comparable data are available."
            ),
        ],
    )


def add_gnet_monero_context(doc: Document) -> None:
    """
    Add qualitative virtual-assets context.
    """

    doc.add_heading("Virtual Assets Context: Privacy Coins and Illicit-Finance Risk", level=1)

    add_paragraphs(
        doc,
        [
            (
                "Qualitative research on privacy coins and illicit-finance risk provides useful context for the "
                "virtual-assets limitations of this country-level model. Some reporting suggests that certain "
                "illicit actors or support networks may show interest in privacy-focused cryptocurrency because "
                "privacy features can complicate tracing compared with more transparent virtual assets."
            ),
            (
                "This context should not be treated as a model variable in the current version. The baseline model "
                "does not observe wallets, transactions, blockchain flows, crowdfunding campaigns, privacy-coin "
                "usage, or online fundraising appeals. It only measures country-level structural vulnerability."
            ),
            (
                "Future versions could add a virtual-assets supplement using official enforcement actions, FIU "
                "reports, national risk assessments, Travel Rule implementation indicators, regulated-sector "
                "typology reports, and country-level crypto-asset supervision indicators."
            ),
        ],
    )


def add_uae_strategy_context(doc: Document) -> None:
    """
    Add UAE AML/CFT/CPF strategy context.
    """

    doc.add_heading("UAE AML/CFT/CPF National Strategy 2024-2027 Context", level=1)

    add_paragraphs(
        doc,
        [
            (
                "The UAE AML/CFT/CPF National Strategy 2024-2027 provides a useful example of how jurisdictions "
                "can continue strengthening AML/CFT frameworks after FATF-related reforms. The strategy emphasizes "
                "risk-based compliance, supervision, national and international coordination, detection and "
                "disruption of illicit financial activity, legal and regulatory updates, and emerging risks."
            ),
            (
                "This source fits the report because the UAE was removed from the FATF grey list in February 2024, "
                "while the national strategy shows continued attention to AML/CFT and counter-proliferation-financing "
                "capabilities. It illustrates why FATF status should be interpreted dynamically rather than as a "
                "static measure of country risk."
            ),
            (
                "The UAE context also supports future data layers related to post-grey-list reform, supervision, "
                "emerging risks, virtual assets, trade-based money laundering, cyber-enabled financial crime, and "
                "national coordination mechanisms."
            ),
        ],
    )


def add_australia_aml_ctf_context(doc: Document) -> None:
    """
    Add Australia AML/CTF Amendment Act context.
    """

    doc.add_heading("Australia AML/CTF Amendment Act 2024 Context", level=1)

    add_paragraphs(
        doc,
        [
            (
                "Australia's AML/CTF Amendment Act 2024 provides useful legal-reform context for the model's "
                "limitations and future improvements. The reform expands AML/CTF coverage, modernizes virtual-asset "
                "and payments regulation, and simplifies parts of the regime to support more flexible financial-crime "
                "prevention."
            ),
            (
                "The reform is relevant because the baseline model does not directly measure domestic statutory "
                "coverage, regulated-sector scope, customer due diligence requirements, value-transfer reporting, "
                "virtual-asset service coverage, or supervisory powers. These legal and supervisory factors can "
                "affect FATF outcomes and AML/CFT effectiveness."
            ),
            (
                "Future versions of the project could add a legal and regulatory supplement that tracks whether "
                "jurisdictions cover real estate professionals, precious metals and stones dealers, lawyers, "
                "conveyancers, accountants, trust and company service providers, virtual-asset services, and other "
                "high-risk sectors."
            ),
        ],
    )


def add_eba_travel_rule_context(doc: Document) -> None:
    """
    Add EBA Travel Rule Guidelines context.
    """

    doc.add_heading("European Banking Authority Travel Rule Guidelines Context", level=1)

    add_paragraphs(
        doc,
        [
            (
                "The European Banking Authority's 2024 Travel Rule Guidelines provide useful regulatory context "
                "for the virtual-asset and funds-transfer limitations of this country-level model. The guidelines "
                "address information requirements for transfers of funds and certain crypto-asset transfers under "
                "Regulation (EU) 2023/1113."
            ),
            (
                "This strengthens the study's methodological distinction between structural country risk and "
                "operational AML/CFT controls. A country may have elevated structural vulnerability, but transaction-level "
                "detection still depends on regulated institutions, crypto-asset service providers, payment service "
                "providers, supervisory expectations, and information-sharing rules."
            ),
            (
                "For future development, the Travel Rule guidance could support a separate virtual-assets and payments "
                "compliance supplement. That supplement could track whether jurisdictions have implemented transfer-originator "
                "and beneficiary-information rules, whether crypto-asset service providers are covered by AML/CFT obligations, "
                "and whether supervisory frameworks address missing or incomplete transfer information."
            ),
        ],
    )


def add_crs_policy_context(doc: Document) -> None:
    """
    Add CRS AML/CFT policy overview context.
    """

    doc.add_heading("Congressional Research Service AML/CFT Policy Overview Context", level=1)

    add_paragraphs(
        doc,
        [
            (
                "The Congressional Research Service overview of U.S. efforts to combat money laundering, terrorist "
                "financing, and other illicit financial threats provides useful policy context for this study. CRS "
                "describes the United States as maintaining a multifaceted AML/CFT and illicit-finance policy regime "
                "involving legal, regulatory, enforcement, and diplomatic tools."
            ),
            (
                "The CRS overview is especially useful because it explains that terrorist financing involves fundraising "
                "through both licit and illicit means and financially sustaining terrorist groups. It also highlights "
                "that illicit financial threats include money laundering, proliferation finance, sanctions evasion, "
                "tax evasion, and other forms of state or nonstate financial facilitation."
            ),
            (
                "The report reinforces the importance of statutory, regulatory, and supervisory controls that this "
                "model does not directly measure, including the Bank Secrecy Act, USA PATRIOT Act CFT provisions, "
                "FinCEN reporting, financial-institution examinations, beneficial ownership policy, real estate and "
                "investment-adviser rulemakings, and digital-asset AML/CFT regulation."
            ),
        ],
    )


def add_gao_context(doc: Document) -> None:
    """
    Add GAO illicit-finance progress assessment context.
    """

    doc.add_heading("GAO Illicit-Finance Progress Assessment Context", level=1)

    add_paragraphs(
        doc,
        [
            (
                "The U.S. Government Accountability Office's 2025 review of federal illicit-finance efforts provides "
                "useful oversight context for this project. GAO examined selected federal roles, responsibilities, "
                "strategies, and efforts related to countering illicit finance activities, including terrorism "
                "financing, money laundering, transnational organized crime, and related financial-crime threats."
            ),
            (
                "GAO found that federal agencies investigate and prosecute illicit-finance activity through multiple "
                "law enforcement and interagency mechanisms. The report discusses the roles of agencies such as the "
                "FBI, DEA, ICE Homeland Security Investigations, Secret Service, IRS-Criminal Investigation, and "
                "U.S. Attorneys' Offices, as well as task-force structures such as OCDETF and the El Dorado Task Force."
            ),
            (
                "The GAO report is also useful for model-governance discussion because it found that selected federal "
                "strategies and efforts to counter illicit finance did not always have clearly defined goals or regular "
                "collection and assessment of performance information. That finding supports future model-governance "
                "tracking and benchmark monitoring."
            ),
        ],
    )


def add_rusi_context(doc: Document) -> None:
    """
    Add RUSI 2025 terrorist-financing reassessment context.
    """

    doc.add_heading("RUSI 2025 Terrorist-Financing Reassessment Context", level=1)

    add_paragraphs(
        doc,
        [
            (
                "RUSI's 2025 assessment of terrorist financing provides useful qualitative context for the limitations "
                "and future-development path of this country-level OSINT model. The paper describes terrorist financing "
                "as an increasingly complex and technology-enabled landscape, with continuing use of traditional channels "
                "such as cash, hawala, and non-profit organization abuse alongside newer channels such as virtual assets, "
                "fintech platforms, instant payments, online crowdfunding, and online gaming."
            ),
            (
                "This source reinforces the difference between structural country vulnerability and operational detection. "
                "The baseline model scores broad country-level conditions such as governance weakness, political instability, "
                "and economic stress. Modern terrorist-financing detection also depends on transaction monitoring, public-private "
                "cooperation, financial intelligence sharing, virtual-asset tracing, off-chain data, platform cooperation, and "
                "investigative capacity."
            ),
            (
                "Future versions could add supplemental indicators for fintech and instant-payment risk, virtual-asset and "
                "stablecoin exposure, online crowdfunding controls, online-gaming vulnerabilities, public-private partnership "
                "maturity, and supervisory guidance for emerging technology risks."
            ),
        ],
    )



def add_charts(doc: Document) -> None:
    """
    Add visual outputs with short chart explanations.
    """

    doc.add_heading("Charts and Visual Outputs", level=1)

    doc.add_paragraph(
        "The following charts summarize the main quantitative outputs from the model. "
        "They support interpretation of the country rankings, predicted FATF grey-list "
        "probabilities, sensitivity analysis, conflict exposure, regional patterns, "
        "current FATF comparison, ACLED recent signals, and threat-finance collection "
        "prioritization layers."
    )

    charts = [
        (
            TOP_15_RISK_CHART_FILE,
            "Figure 1. Top 15 Countries by Terrorism-Financing Vulnerability Score",
            (
                "This chart shows the countries with the highest structural vulnerability scores. "
                "Higher scores reflect a stronger combination of governance weakness, political instability, "
                "lower income conditions, inflation stress, and unemployment stress under the model's weighting framework."
            ),
        ),
        (
            RISK_SCORE_PROBABILITY_CHART_FILE,
            "Figure 2. Risk Score vs Predicted FATF Grey-List Probability",
            (
                "This scatter plot compares the weighted vulnerability score with the model's predicted probability "
                "of FATF grey-list status. A positive relationship suggests that the selected public indicators provide "
                "screening value, although the probability should not be treated as a definitive compliance or risk finding."
            ),
        ),
        (
            GREY_LIST_PROBABILITY_TOP_20_CHART_FILE,
            "Figure 3. Top 20 Countries by Predicted FATF Grey-List Probability",
            (
                "This chart ranks countries by the classification model's estimated probability of FATF grey-list status. "
                "Countries appearing here may warrant deeper qualitative review, especially when model probability is high "
                "but observed FATF status differs."
            ),
        ),
        (
            RISK_TIER_DISTRIBUTION_CHART_FILE,
            "Figure 4. Risk Tier Distribution",
            (
                "This chart shows how countries are distributed across the project's risk tiers. It translates the "
                "continuous 0-to-100 score into easier-to-communicate categories such as Very High, High, Moderate, Low, "
                "and Very Low."
            ),
        ),
        (
            SENSITIVITY_OVERLAP_CHART_FILE,
            "Figure 5. Top-20 Sensitivity Overlap by Scenario",
            (
                "This chart tests ranking stability across alternative weighting assumptions. Higher overlap means the "
                "same countries remain in the top-risk group even when the model places more emphasis on governance, "
                "political instability, or economic stress."
            ),
        ),
        (
            TOP_15_CONFLICT_FATALITIES_CHART_FILE,
            "Figure 6. Top 15 Countries by Recorded Organized-Violence Deaths",
            (
                "This chart summarizes supplemental conflict exposure using UCDP organized-violence data. Conflict deaths "
                "are not part of the baseline vulnerability score, but they provide useful instability context for interpreting "
                "country-level AML/CFT risk environments."
            ),
        ),
        (
            REGIONAL_AVERAGE_RISK_CHART_FILE,
            "Figure 7. Regional Average Terrorism-Financing Vulnerability Score",
            (
                "This chart compares average vulnerability scores across World Bank regional groupings. It helps identify "
                "whether elevated country-level risk is concentrated in specific regions and provides context for regional "
                "screening priorities."
            ),
        ),
        (
            REGIONAL_FATF_COUNT_CHART_FILE,
            "Figure 8. Current FATF Increased-Monitoring Count by Region",
            (
                "This chart shows how current FATF increased-monitoring jurisdictions are distributed by region. It is a "
                "current-status comparison layer and should be interpreted separately from the baseline quantitative score."
            ),
        ),
        (
            CURRENT_FATF_STATUS_COMPARISON_CHART_FILE,
            "Figure 9. Current FATF Status Comparison",
            (
                "This chart compares the model's FATF target coding with the current FATF increased-monitoring status "
                "configured in the project. It helps show where the model-year target and current-status supplement align "
                "or differ."
            ),
        ),
        (
            TOP_15_COLLECTION_PRIORITY_CHART_FILE,
            "Figure 10. Top 15 Threat Finance Collection Priorities",
            (
                "This chart summarizes the countries with the highest collection-priority scores. The score combines "
                "baseline vulnerability, live or recent signal strength, intelligence gaps, source staleness, and "
                "decision relevance. It is a public-source review prioritization output, not a terrorism-financing finding."
            ),
        ),
        (
            INTELLIGENCE_GAP_DISTRIBUTION_CHART_FILE,
            "Figure 11. Distribution of Threat Finance Intelligence Gap Scores",
            (
                "This chart shows how intelligence-gap scores are distributed across the modeled countries. Higher values "
                "indicate more missing, stale, or weak source coverage and can help analysts identify where refreshed "
                "public-source collection is most needed."
            ),
        ),
        (
            ACLED_SIGNAL_TOP_15_CHART_FILE,
            "Figure 12. Top 15 ACLED Recent Conflict and Unrest Signals",
            (
                "This chart summarizes recent ACLED conflict, unrest, and political-violence signals. It supports the "
                "live-signal layer and collection-priority workflow, but it does not identify terrorism financing or "
                "criminal activity."
            ),
        ),
        (
            TOP_15_ORGANIZED_CRIME_CHART_FILE,
            "Figure 13. Top 15 Organized Crime Vulnerability Scores",
            (
                "This chart summarizes supplemental Global Organized Crime Index enrichment. It highlights where "
                "organized-crime exposure and resilience gaps may provide useful convergence context for analyst review. "
                "It is not evidence of terrorism financing, criminality, sanctions exposure, or illicit activity."
            ),
        ),
        (
            TOP_15_GDELT_NEWS_SIGNAL_CHART_FILE,
            "Figure 14. Top 15 GDELT 2.0 Batch Public-News Signal Scores",
            (
                "This chart summarizes supplemental GDELT 2.0 batch public-news signal enrichment. It highlights where recent "
                "media-reporting volume may support refreshed public-source review. It is not evidence of terrorism "
                "financing, criminality, sanctions exposure, or illicit activity."
            ),
        ),
    ]

    for chart_path, caption, explanation in charts:
        add_picture_if_exists(doc, chart_path, caption)
        doc.add_paragraph(explanation)



def add_correlation_analysis(doc: Document, corr_df: pd.DataFrame | None) -> None:
    """
    Add correlation analysis.
    """

    doc.add_heading("Correlation Analysis", level=1)

    doc.add_paragraph(
        "The correlation matrix provides a basic view of how the model's risk indicators relate to the final "
        "vulnerability score, FATF grey-list status, and predicted grey-list probability."
    )

    if corr_df is None or corr_df.empty:
        doc.add_paragraph("Correlation matrix was not available.")
        return

    display_df = prepare_correlation_table(corr_df)

    add_dataframe_table(
        doc,
        display_df,
        ["Variable", "Risk Score", "Grey List", "Pred. Prob."],
        ["Variable", "Risk Score", "Grey List", "Pred. Prob."],
    )


def add_limitations(doc: Document) -> None:
    """
    Add limitations.
    """

    doc.add_heading("Limitations", level=1)

    limitations = [
        "The model uses country-level indicators and does not identify individual actors, entities, or transactions.",
        "FATF grey-list status is a policy and regulatory outcome, not a direct measure of terrorism financing activity.",
        "The model does not include detailed legal, supervisory, enforcement, or financial intelligence data.",
        "Some countries may have missing or delayed economic data, so the pipeline records the latest available observation year used for each indicator.",
        "The project does not use 2025 or 2026 as the primary model year because the core public country-level indicators used for scoring are not yet available as a complete and comparable 2025 or 2026 panel.",
        "Conflict event data is supplemental context and is not evidence of terrorist-financing activity.",
        "UCDP country matching depends on country-name harmonization; unmatched countries may require a manual crosswalk.",
        "Sanctions exposure and FATF mutual-evaluation supplements are currently template-based unless structured source values are manually populated.",
        "The model is intended for defensive screening and academic analysis, not operational targeting.",
        "The model does not replace suspicious activity reporting, law enforcement investigations, or the multi-indicator assessment process described in official counterterrorism finance guidance.",
        "The model is not a substitute for institution-specific AML/CFT risk assessment processes across products, services, customers, distribution channels, and geographic locations.",
        "The model does not capture all current threat-specific channels discussed in national risk assessments, such as online fundraising, money services businesses, cash, digital assets, or nonprofit organization vulnerabilities.",
        "The model does not capture FATF's full terrorist-financing risk indicator framework, including detailed customer behavior, transaction, geographic, product/service, trade, commercial-entity, nonprofit, and organized-crime indicators.",
        "The model does not replace bank-level AML/CFT risk management practices, including customer due diligence, ongoing monitoring, suspicious-transaction reporting, governance arrangements, and the three lines of defense.",
        "The model does not directly measure virtual-asset activity, privacy-coin exposure, wallet-level behavior, crowdfunding, or online fundraising appeals.",
        "The model does not directly measure Travel Rule compliance, originator and beneficiary information quality, missing or incomplete transfer information, payment-service-provider controls, or crypto-asset-service-provider controls.",
        "The model does not directly measure domestic AML/CFT statutory coverage, beneficial ownership implementation, FinCEN-style reporting systems, supervisory examinations, enforcement actions, or congressional policy developments.",
        "The model does not measure interagency coordination, investigative case outcomes, prosecution results, asset forfeiture, task-force disruption metrics, or government-wide performance tracking for illicit-finance strategies.",
        "The model does not directly measure fintech-platform controls, instant-payment exposure, online crowdfunding activity, online-gaming financial flows, public-private partnership maturity, or cross-sector intelligence-sharing capacity.",
        "The current data supplement is a comparison layer and should not be treated as definitive proof of current terrorist-financing activity or absence of risk.",
    ]

    for limitation in limitations:
        doc.add_paragraph(limitation, style="List Bullet")


def add_conclusion(doc: Document) -> None:
    """
    Add conclusion and future improvements.
    """

    doc.add_heading("Conclusion", level=1)

    add_paragraphs(
        doc,
        [
            (
                "This project demonstrates how open-source data can be used to build a structured, defensive "
                "intelligence model for country-level terrorism financing vulnerability. The model combines "
                "governance weakness, political instability, and economic stress into a transparent scoring "
                "framework and compares the results against FATF grey-list status."
            ),
            (
                "The expanded project adds conflict exposure, regional summaries, top risk drivers, additional "
                "charts, official AML/CFT policy context, virtual-assets context, Travel Rule context, legal and "
                "regulatory reform examples, and structured input templates for sanctions and FATF mutual-evaluation data."
            ),
            (
                "The results suggest that open-source indicators can provide useful screening value, but they should "
                "be supplemented with qualitative analysis, AML/CFT evaluation reports, sanctions data, conflict event "
                "data, official risk assessments, and expert judgment for higher-confidence assessment."
            ),
        ],
    )

    doc.add_heading("Future Improvements", level=2)

    future_items = [
        "Add sanctions exposure indicators with populated source data.",
        "Add a virtual-assets and payments compliance supplement using Travel Rule, payment-transfer, and crypto-asset transfer implementation indicators.",
        "Add a legal and regulatory framework supplement covering AML/CFT statutory scope, beneficial ownership transparency, supervisory coverage, enforcement activity, and digital-asset AML/CFT regulation.",
        "Add an emerging terrorist-financing typologies supplement covering fintech, instant payments, virtual assets, online crowdfunding, online gaming, and public-private information-sharing maturity.",
        "Add model-governance tracking that evaluates whether supplemental indicators improve screening performance, ranking stability, and alignment with official illicit-finance strategy benchmarks over time.",
        "Refine conflict event data matching with a country-name crosswalk.",
        "Add AML/CFT mutual evaluation data with technical compliance and effectiveness ratings.",
        "Expand sensitivity analysis with additional weighting scenarios.",
        "Refine risk-tier interpretation with regional and peer-group comparisons.",
        "Add additional current FATF comparison visuals over multiple FATF update cycles.",
        "Build an interactive Streamlit dashboard.",
        "Add time-series analysis across multiple years.",
        "Compare results across FATF grey-list updates.",
    ]

    for item in future_items:
        doc.add_paragraph(item, style="List Bullet")


def add_references(doc: Document) -> None:
    """
    Add references and data sources.
    """

    doc.add_heading("References and Data Sources", level=1)

    doc.add_paragraph(
        "The following sources were used to support the data collection, methodology, validation framework, "
        "limitations, and supplemental context for this project."
    )

    references = [
        "World Bank. Worldwide Governance Indicators. Used for Control of Corruption, Rule of Law, and Political Stability indicators. https://www.worldbank.org/en/publication/worldwide-governance-indicators",
        "World Bank. WGI 2025 Revision / WGI 2.0. Used to document current WGI coverage and updated governance-indicator methodology context. https://www.govindicators.org/",
        "World Bank DataBank. Worldwide Governance Indicators. Used as an official data access source for WGI indicators. https://databank.worldbank.org/source/worldwide-governance-indicators",
        "World Bank API / World Development Indicators. Used for GDP per capita, inflation, unemployment, and country metadata indicators. https://api.worldbank.org/v2/country/all/indicator",
        "Uppsala Conflict Data Program. UCDP Country-Year Dataset on Organized Violence within Country Borders. Used as supplemental country-year conflict exposure context. https://ucdp.uu.se/downloads/",
        "Financial Action Task Force. FATF Home. Used as the official institutional source for FATF's AML/CFT role. https://www.fatf-gafi.org/en/home.html",
        "Financial Action Task Force. Black and Grey Lists. Used for FATF grey-list status and jurisdiction monitoring context. https://www.fatf-gafi.org/en/countries/black-and-grey-lists.html",
        "Financial Action Task Force. FATF Recommendations. Used for AML/CFT standards and policy framework context. https://www.fatf-gafi.org/content/dam/fatf-gafi/recommendations/FATF%20Recommendations%202012.pdf.coredownload.inline.pdf",
        "Financial Action Task Force. Comprehensive Update on Terrorist Financing Risks. July 2025. Used as current FATF context for evolving terrorist-financing risk factors, methods, indicators, and recommendations. https://www.fatf-gafi.org/content/dam/fatf-gafi/publications/Comprehensive-Update-on-Terrorist-Financing-Risks-2025.pdf.coredownload.inline.pdf",
        "U.S. Department of the Treasury. National Terrorist Financing Risk Assessment. Used as government context for terrorist-financing threats, vulnerabilities, financial channels, and risk-assessment methodology. https://home.treasury.gov/policy-issues/terrorism-and-illicit-finance/national-risk-assessments",
        "U.S. Department of the Treasury. 2024 National Strategy for Combating Terrorist and Other Illicit Financing. Used as strategic context for U.S. illicit-finance risk prioritization, AML/CFT modernization, supervision, transparency, enforcement, and disruption of illicit-finance networks. https://home.treasury.gov/system/files/136/2024-Illicit-Finance-Strategy.pdf",
        "National Counterterrorism Center, Department of Homeland Security, and Federal Bureau of Investigation. Identifying and Preventing Terrorist and Other Illicit Financing. Used as official U.S. government context for terrorist and illicit financing indicators, investigative considerations, and public-private collaboration. https://www.dni.gov/files/NCTC/documents/jcat/firstresponderstoolbox/135s_-_First_Responders_Toolbox_-_Identifying_and_Preventing_Terrorist_and_Other_Illicit_Financing.pdf",
        "INTERPOL. Global Counter-Terrorism Strategy: Summary. Used as international law-enforcement context for counter-terrorism action streams, including finances, intelligence analysis, police data management, and capacity building. https://www.interpol.int",
        "Basel Committee on Banking Supervision. Sound Management of Risks Related to Money Laundering and Financing of Terrorism. Bank for International Settlements, June 2017. Used as banking-supervision context for ML/FT risk management, customer due diligence, governance, monitoring, suspicious-transaction reporting, and cross-border AML/CFT controls. https://www.bis.org/bcbs/publ/d405.pdf",
        "Financial Intelligence Centre. Annual Report 2024/25. Used as South Africa FIU and FATF grey-list remediation case-study context.",
        "United Arab Emirates. AML/CFT/CPF National Strategy 2024-2027. Used as post-grey-list reform and national-strategy context.",
        "European Banking Authority. Travel Rule Guidelines. Used as virtual-assets and funds-transfer compliance context.",
        "Congressional Research Service. AML/CFT policy overview. Used as U.S. policy and statutory context.",
        "U.S. Government Accountability Office. Illicit-finance progress assessment. Used as model-governance and oversight context.",
        "Royal United Services Institute. Terrorist-financing reassessment context. Used as qualitative current-risk and future-improvement context.",
    ]

    for reference in references:
        doc.add_paragraph(reference, style="List Number")


def add_logit_appendix(doc: Document, logit_text: str) -> None:
    """
    Add the logistic regression output as an appendix.
    """

    doc.add_heading("Appendix: Logistic Regression Output", level=1)

    if not logit_text:
        doc.add_paragraph("Logistic regression summary file was not available.")
        return

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)

    run = paragraph.add_run(logit_text)
    run.font.name = "Courier New"
    run.font.size = Pt(7)


# ---------------------------------------------------
# Report generation
# ---------------------------------------------------

def build_report() -> Document:
    """
    Build the full Word report document.
    """

    rankings_df = read_csv_if_exists(RISK_RANKINGS_FILE)
    corr_df = read_csv_if_exists(CORRELATION_MATRIX_FILE)
    risk_tier_df = read_csv_if_exists(RISK_TIER_SUMMARY_FILE)
    performance_df = read_csv_if_exists(MODEL_PERFORMANCE_FILE)
    false_positive_df = read_csv_if_exists(FALSE_POSITIVES_FILE)
    false_negative_df = read_csv_if_exists(FALSE_NEGATIVES_FILE)
    sensitivity_overlap_df = read_csv_if_exists(SENSITIVITY_OVERLAP_FILE)
    current_df = read_csv_if_exists(CURRENT_DATA_SUPPLEMENT_FILE)
    current_fatf_df = read_csv_if_exists(CURRENT_FATF_STATUS_COMPARISON_FILE)
    conflict_df = read_csv_if_exists(CONFLICT_DATA_SUPPLEMENT_FILE)
    top_drivers_df = read_csv_if_exists(TOP_RISK_DRIVERS_FILE)
    regional_df = read_csv_if_exists(REGIONAL_RISK_SUMMARY_FILE)
    summary_df = read_csv_if_exists(ADDITIONAL_RISK_FACTOR_SUMMARY_FILE)
    sanctions_df = read_csv_if_exists(SANCTIONS_EXPOSURE_SUPPLEMENT_FILE)
    mutual_eval_df = read_csv_if_exists(FATF_MUTUAL_EVALUATION_SUPPLEMENT_FILE)
    collection_priority_df = read_csv_if_exists(
        THREAT_FINANCE_COLLECTION_PRIORITIES_FILE
    )
    source_register_df = read_csv_if_exists(
        THREAT_FINANCE_SOURCE_REGISTER_FILE
    )
    acled_signal_df = read_csv_if_exists(ACLED_SIGNAL_ENRICHMENT_FILE)
    organized_crime_df = read_csv_if_exists(ORGANIZED_CRIME_ENRICHMENT_FILE)
    gdelt_news_df = read_csv_if_exists(GDELT_NEWS_SIGNAL_ENRICHMENT_FILE)
    logit_text = read_text_if_exists(LOGIT_SUMMARY_FILE)

    doc = Document()
    set_document_style(doc)

    add_title_page(doc)
    add_executive_summary(doc, rankings_df)
    add_project_overview(doc)
    add_methodology(doc, rankings_df)
    add_top_country_rankings(doc, rankings_df)
    add_fatf_grey_list_background(doc)
    add_modeling_results(doc, rankings_df)
    add_enhanced_analysis(
        doc,
        risk_tier_df,
        performance_df,
        false_positive_df,
        false_negative_df,
        sensitivity_overlap_df,
    )
    add_top_risk_drivers(doc, top_drivers_df)
    add_conflict_supplement(doc, conflict_df)
    add_additional_risk_factors(
        doc,
        regional_df,
        summary_df,
        sanctions_df,
        mutual_eval_df,
    )
    add_threat_finance_intelligence_gap_analysis(
        doc,
        collection_priority_df,
        source_register_df,
    )
    add_acled_signal_supplement(doc, acled_signal_df)
    add_organized_crime_supplement(doc, organized_crime_df)
    add_gdelt_news_signal_supplement(doc, gdelt_news_df)
    add_current_data_supplement(doc, current_df, current_fatf_df)
    add_official_designation_context(doc)
    add_south_africa_fic_context(doc)
    add_treasury_ntfra_context(doc)
    add_treasury_strategy_context(doc)
    add_gnet_monero_context(doc)
    add_uae_strategy_context(doc)
    add_australia_aml_ctf_context(doc)
    add_eba_travel_rule_context(doc)
    add_crs_policy_context(doc)
    add_gao_context(doc)
    add_rusi_context(doc)
    add_charts(doc)
    add_correlation_analysis(doc, corr_df)
    add_limitations(doc)
    add_conclusion(doc)
    add_references(doc)
    add_logit_appendix(doc, logit_text)

    return doc


def save_word_report(doc: Document) -> None:
    """
    Save the Word report.
    """

    WORD_REPORT_FILE.parent.mkdir(parents=True, exist_ok=True)
    doc.save(WORD_REPORT_FILE)

    print("Word report saved to:")
    print(WORD_REPORT_FILE)


def save_pdf_report() -> None:
    """
    Convert the Word report to PDF using docx2pdf when available.

    On Windows, docx2pdf uses Microsoft Word. If conversion is unavailable,
    the Word report is still saved successfully.
    """

    try:
        from docx2pdf import convert

        convert(str(WORD_REPORT_FILE), str(PDF_REPORT_FILE))

        print("PDF report saved to:")
        print(PDF_REPORT_FILE)

    except Exception as error:
        print("PDF conversion was skipped or failed.")
        print(f"Reason: {error}")
        print("The Word report was still created successfully.")


def main() -> None:
    """
    Generate the full report.
    """

    print("=" * 70)
    print("Generating report")
    print("=" * 70)

    doc = build_report()
    save_word_report(doc)
    save_pdf_report()

    print("\nReport generation complete.")


if __name__ == "__main__":
    main()